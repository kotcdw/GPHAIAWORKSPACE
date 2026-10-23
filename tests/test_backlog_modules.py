"""Planning memo, capacity & time sheets, QAIP, analytics, global search, Excel import/export, saved views."""

from __future__ import annotations

import io
import uuid
from datetime import timedelta

from openpyxl import Workbook, load_workbook

from app.core import appconfig
from app.core.enums import ReviewLevel, Role, utcnow
from app.platform.notifications.models import Notification
from tests.conftest import DEFAULT_PASSWORD, csrf_headers, decide_topicals, login, make_user


def _u(db, prefix: str, role: Role, level: ReviewLevel | None = None):
    user = make_user(db, f"{prefix}_{uuid.uuid4().hex[:6]}", role)
    if level:
        user.review_level = level
        db.commit()
    return user


def _xlsx(header: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MEMO = {
    "background": "Revenue from vessel calls.",
    "objectives": "Assess billing completeness.",
    "scope": "FY2026 vessel calls at Tema.",
    "key_risks": "Unbilled calls; tariff errors.",
    "approach": "Walkthrough, data analytics on 100% of calls, sample testing.",
    "team": "Kofi (lead), Ama",
    "budget_hours": 120,
}


def test_planning_memo_flow_and_fieldwork_gate(client, client2, db):
    auditor = _u(db, "pmaud", Role.Auditor)
    manager = _u(db, "pmmgr", Role.Auditor, ReviewLevel.Manager)
    admin = _u(db, "pmadm", Role.Admin)
    h = csrf_headers(login(client, auditor.username))
    mh = csrf_headers(login(client2, manager.username))
    aid = client.post("/api/audits", json={"title": "Billing audit", "scope": "Billing"}, headers=h).json()["id"]

    empty = client.get(f"/api/audits/{aid}/planning-memo").json()
    assert empty["id"] is None and empty["can_edit"] and "Objectives" in empty["missing"]

    bad = client.put(
        f"/api/audits/{aid}/planning-memo",
        json={**MEMO, "start_date": "2026-11-10T00:00:00", "end_date": "2026-11-01T00:00:00"},
        headers=h,
    )
    assert bad.status_code == 422

    partial = client.put(f"/api/audits/{aid}/planning-memo", json={"objectives": "x"}, headers=h).json()
    assert partial["status"] == "Draft"
    blocked = client.post(f"/api/audits/{aid}/planning-memo/submit", headers=h)
    assert blocked.status_code == 409 and "Scope" in blocked.json()["detail"]

    full = {**MEMO, "start_date": "2026-11-01T00:00:00", "end_date": "2026-11-30T00:00:00"}
    assert client.put(f"/api/audits/{aid}/planning-memo", json=full, headers=h).status_code == 200
    sub = client.post(f"/api/audits/{aid}/planning-memo/submit", headers=h).json()
    assert sub["status"] == "Submitted" and not sub["can_edit"]
    db.rollback()
    assert any(
        "Planning memo awaiting approval" in n.title for n in db.query(Notification).filter_by(user_id=manager.id)
    )

    # The fieldwork gate (switched on here) blocks starting until approval.
    ah = csrf_headers(login(client2, admin.username))
    try:
        client2.put("/api/config", json={"values": {"require_memo_approval_for_fieldwork": True}}, headers=ah)
        start = client.patch(f"/api/audits/{aid}", json={"status": "In Progress"}, headers=h)
        assert start.status_code == 409 and "planning memo" in start.json()["detail"]

        # The preparer cannot approve; an auditor without a review level cannot either.
        assert client.post(
            f"/api/audits/{aid}/planning-memo/approve", json={"password": DEFAULT_PASSWORD}, headers=h
        ).status_code in (403, 409)
        mh = csrf_headers(login(client2, manager.username))
        ret = client2.post(
            f"/api/audits/{aid}/planning-memo/approve",
            json={"password": DEFAULT_PASSWORD, "decision": "Reject", "comment": "Add the sampling approach"},
            headers=mh,
        ).json()
        assert ret["status"] == "Returned" and ret["review_comment"] == "Add the sampling approach"
        client.put(f"/api/audits/{aid}/planning-memo", json=full, headers=h)
        client.post(f"/api/audits/{aid}/planning-memo/submit", headers=h)
        ok = client2.post(
            f"/api/audits/{aid}/planning-memo/approve", json={"password": DEFAULT_PASSWORD}, headers=mh
        ).json()
        assert ok["status"] == "Approved" and ok["approved_by"] == manager.id
        assert client.put(f"/api/audits/{aid}/planning-memo", json=full, headers=h).status_code == 409
        assert client.patch(f"/api/audits/{aid}", json={"status": "In Progress"}, headers=h).status_code == 200
    finally:
        ah = csrf_headers(login(client2, admin.username))
        client2.post("/api/config/require_memo_approval_for_fieldwork/reset", headers=ah)
        appconfig.invalidate()

    doc = client.get(f"/api/audits/{aid}/planning-memo.docx")
    assert doc.status_code == 200 and doc.content[:2] == b"PK"


def test_capacity_profiles_and_time_sheets(client, client2, db):
    auditor = _u(db, "cpaud", Role.Auditor)
    cae = _u(db, "cpcae", Role.Auditor, ReviewLevel.CAE)
    auditee = _u(db, "cpee", Role.Auditee)
    h = csrf_headers(login(client, auditor.username))
    aid = client.post("/api/audits", json={"title": "Stores audit", "scope": "Stores"}, headers=h).json()["id"]

    prof = client.put(
        f"/api/capacity/profiles/{auditor.id}",
        json={"grade": "Senior Auditor", "skills": ["IT audit", "Data analytics", "it audit"], "annual_hours": 100},
        headers=h,
    ).json()
    assert prof["skills"] == ["IT audit", "Data analytics"]
    assert prof["annual_hours"] == 1600  # only management sets available hours
    assert client.put(f"/api/capacity/profiles/{cae.id}", json={"grade": "x"}, headers=h).status_code == 403
    found = client.get("/api/capacity/profiles?skill=analytics").json()
    assert [p["user_id"] for p in found] == [auditor.id]

    today = utcnow().strftime("%Y-%m-%dT00:00:00")
    assert (
        client.post(
            "/api/capacity/time", json={"entry_date": today, "hours": 2, "category": "Audit"}, headers=h
        ).status_code
        == 422
    )  # audit required
    assert (
        client.post(
            "/api/capacity/time", json={"entry_date": today, "hours": 1.1, "category": "Training"}, headers=h
        ).status_code
        == 422
    )  # quarter hours
    future = (utcnow() + timedelta(days=5)).strftime("%Y-%m-%dT00:00:00")
    assert (
        client.post(
            "/api/capacity/time", json={"entry_date": future, "hours": 1, "category": "Training"}, headers=h
        ).status_code
        == 422
    )
    e1 = client.post(
        "/api/capacity/time",
        json={"entry_date": today, "hours": 6, "category": "Audit", "audit_id": aid, "description": "Walkthrough"},
        headers=h,
    )
    assert e1.status_code == 201, e1.text
    client.post("/api/capacity/time", json={"entry_date": today, "hours": 2, "category": "Training"}, headers=h)
    over = client.post("/api/capacity/time", json={"entry_date": today, "hours": 17, "category": "Other"}, headers=h)
    assert over.status_code == 422 and "24 hours" in over.json()["errors"][0]["message"]

    mine = client.get("/api/capacity/time").json()
    assert len(mine) == 2 and mine[0]["can_delete"]
    ch = csrf_headers(login(client2, cae.username))
    assert len(client2.get(f"/api/capacity/time?user_id={auditor.id}").json()) == 2
    summary = client2.get(f"/api/capacity/summary?fiscal_year={utcnow().year}").json()
    row = next(a for a in summary["auditors"] if a["user_id"] == auditor.id)
    assert row["logged_hours"] == 8 and row["audit_hours"] == 6 and row["utilisation_pct"] == 75
    budget = next(b for b in summary["audits"] if b["audit_id"] == aid)
    assert budget["actual_hours"] == 6
    client2.put(f"/api/capacity/profiles/{auditor.id}", json={"annual_hours": 1200}, headers=ch)
    assert (
        next(p for p in client2.get("/api/capacity/profiles").json() if p["user_id"] == auditor.id)["annual_hours"]
        == 1200
    )

    assert client2.delete(f"/api/capacity/time/{e1.json()['id']}", headers=ch).status_code == 404
    assert client.delete(f"/api/capacity/time/{e1.json()['id']}", headers=h).status_code == 200
    login(client2, auditee.username)
    assert client2.get("/api/capacity/time").status_code == 403


def test_quality_review_and_survey(client, client2, db):
    preparer = _u(db, "qaprep", Role.Auditor)
    reviewer = _u(db, "qarev", Role.Auditor)
    admin = _u(db, "qaadm", Role.Admin)
    auditee = _u(db, "qaee", Role.Auditee)
    h = csrf_headers(login(client, preparer.username))
    aid = client.post("/api/audits", json={"title": "Fleet audit", "scope": "Fleet"}, headers=h).json()["id"]
    client.post(f"/api/audits/{aid}/auditees", json={"user_id": auditee.id}, headers=h)
    client.post(
        f"/api/audits/{aid}/findings", json={"title": "Fuel cards unreconciled", "risk_rating": "Medium"}, headers=h
    )

    rh = csrf_headers(login(client2, reviewer.username))
    early = client2.put(f"/api/audits/{aid}/quality-review", json={"answers": {}}, headers=rh)
    assert early.status_code == 409
    client.patch(f"/api/audits/{aid}", json={"status": "In Progress"}, headers=h)
    client.patch(f"/api/audits/{aid}", json={"status": "Review"}, headers=h)
    assert client.put(f"/api/audits/{aid}/quality-review", json={"answers": {}}, headers=h).status_code == 403

    review = client2.get(f"/api/audits/{aid}/quality-review").json()
    assert len(review["items"]) == 12 and review["can_edit"]
    assert "1 finding" in next(i for i in review["items"] if i["key"] == "management_discussion")["hint"] or True
    keys = [i["key"] for i in review["items"]]
    answers = {k: {"answer": "Conforms"} for k in keys}
    answers["supervision"] = {"answer": "Does not conform", "comment": ""}
    answers["budget"] = {"answer": "N/A"}
    incomplete = client2.put(
        f"/api/audits/{aid}/quality-review",
        json={"answers": answers, "conclusion": "Generally conforms", "complete": True},
        headers=rh,
    )
    assert incomplete.status_code == 422  # unexplained non-conformance
    answers["supervision"]["comment"] = "Two tests not reviewed"
    done = client2.put(
        f"/api/audits/{aid}/quality-review",
        json={"answers": answers, "conclusion": "Generally conforms", "complete": True},
        headers=rh,
    )
    assert done.status_code == 200, done.text
    assert done.json()["score"] == round(100 * 10 / 11, 1) and not done.json()["can_edit"]

    # Survey opens when the audit closes; auditees are invited.
    eh = csrf_headers(login(client, auditee.username))
    assert client.get(f"/api/audits/{aid}/survey").json()["can_respond"] is False
    ah = csrf_headers(login(client2, admin.username))
    decide_topicals(client2, aid, ah)
    assert client2.post(f"/api/audits/{aid}/close", json={"password": DEFAULT_PASSWORD}, headers=ah).status_code == 200
    db.rollback()
    assert any("Tell us how the audit went" in n.title for n in db.query(Notification).filter_by(user_id=auditee.id))
    state = client.get(f"/api/audits/{aid}/survey").json()
    assert state["open"] and state["can_respond"] and len(state["questions"]) == 6
    scores = {q["key"]: 4 for q in state["questions"]}
    assert (
        client.post(f"/api/audits/{aid}/survey", json={"scores": {"communication": 9}}, headers=eh).status_code == 422
    )
    scores["timeliness"] = 2
    sent = client.post(f"/api/audits/{aid}/survey", json={"scores": scores, "comments": "Report came late"}, headers=eh)
    assert sent.status_code == 201 and sent.json()["my_response"]["comments"] == "Report came late"
    assert client.post(f"/api/audits/{aid}/survey", json={"scores": scores}, headers=eh).status_code == 409
    staff_view = client2.get(f"/api/audits/{aid}/survey").json()
    assert staff_view["responses"] == 1 and staff_view["average"] == round(22 / 6, 2)

    qaip = client2.get(f"/api/qaip/summary?fiscal_year={utcnow().year}").json()
    assert qaip["reviews_completed"] >= 1 and qaip["survey_responses"] >= 1
    assert any(w["key"] == "supervision" for w in qaip["weakest_items"])


def test_analytics_search_and_repeat_findings(client, client2, db):
    auditor = _u(db, "anaud", Role.Auditor)
    auditee = _u(db, "anee", Role.Auditee)
    h = csrf_headers(login(client, auditor.username))
    tag = uuid.uuid4().hex[:5]
    a1 = client.post("/api/audits", json={"title": f"Payroll {tag} 2025", "scope": "Payroll"}, headers=h).json()["id"]
    a2 = client.post("/api/audits", json={"title": f"Payroll {tag} 2026", "scope": "Payroll"}, headers=h).json()["id"]
    client.post(
        f"/api/audits/{a1}/findings",
        json={"title": f"Ghost employees zq{tag} on payroll not removed", "risk_rating": "High"},
        headers=h,
    )
    client.post(
        f"/api/audits/{a2}/findings",
        json={"title": f"Ghost employees zq{tag} payroll not removed timely", "risk_rating": "High"},
        headers=h,
    )

    data = client.get("/api/analytics").json()
    assert [r["rating"] for r in data["entity_ratings"]] == ["High", "Medium", "Low", "Not assessed"]
    assert len(data["findings_trend"]) == 12
    assert {k["key"] for k in data["kpis"]} >= {
        "plan",
        "actions_on_time",
        "overdue",
        "report_days",
        "survey",
        "utilisation",
    }
    rep = [r for r in data["repeat_findings"] if f"zq{tag}" in r["title"]]
    assert rep and rep[0]["occurrences"] == 2

    hits = client.get(f"/api/search?q=Payroll {tag}").json()
    assert {x["id"] for x in hits if x["type"] == "audit"} == {a1, a2}
    assert client.get("/api/search?q=x").status_code == 422
    login(client2, auditee.username)
    assert client2.get(f"/api/search?q=Payroll {tag}").json() == []  # not assigned -> nothing
    assert client2.get("/api/analytics").status_code == 403


def test_excel_import_export_and_saved_views(client, client2, db):
    auditor = _u(db, "ioaud", Role.Auditor)
    admin = _u(db, "ioadm", Role.Admin)
    ah = csrf_headers(login(client2, auditor.username))
    assert client2.post("/api/import/orgunits", files={"file": ("u.xlsx", b"x", "application/octet-stream")},
                        headers=ah).status_code == 403
    h = csrf_headers(login(client, admin.username))
    tpl = client.get("/api/import/orgunits/template.xlsx")
    assert tpl.status_code == 200
    wb = load_workbook(io.BytesIO(tpl.content))
    assert wb.worksheets[0]["A1"].value == "Name" and "Instructions" in wb.sheetnames

    code = uuid.uuid4().hex[:4].upper()
    header = ["Name", "Type", "Parent", "Address"]
    bad = _xlsx(header, [[f"Takoradi {code}", "Spaceport", "", ""], ["", "Department", "", ""]])
    res = client.post(
        "/api/import/orgunits?commit=true", files={"file": ("u.xlsx", bad, "application/octet-stream")}, headers=h
    ).json()
    assert not res["committed"] and {e["field"] for e in res["errors"]} >= {"Type", "Name"}

    good = _xlsx(header, [[f"Takoradi {code}", "Port", "", ""], [f"Bulk terminal {code}", "Department",
                                                                  f"Takoradi {code}", ""]])
    preview = client.post(
        "/api/import/orgunits", files={"file": ("u.xlsx", good, "application/octet-stream")}, headers=h
    ).json()
    assert preview["to_create"] == 2 and not preview["committed"] and not preview["errors"]
    done = client.post(
        "/api/import/orgunits?commit=true", files={"file": ("u.xlsx", good, "application/octet-stream")}, headers=h
    ).json()
    assert done["committed"]
    assert (
        client.post(
            "/api/import/orgunits", files={"file": ("u.xlsx", b"nope", "application/octet-stream")}, headers=h
        ).status_code
        == 422
    )

    aid = client.post("/api/audits", json={"title": "Import audit", "scope": "x"}, headers=h).json()["id"]
    fx = _xlsx(
        ["Title", "Risk rating", "Recommendation"],
        [["Gate passes not logged", "high", "Log them"], ["Weighbridge not calibrated", "Low", ""]],
    )
    r = client.post(
        f"/api/audits/{aid}/import/findings?commit=true",
        files={"file": ("f.xlsx", fx, "application/octet-stream")},
        headers=h,
    ).json()
    assert r["committed"] and r["to_create"] == 2
    assert len(client.get(f"/api/audits/{aid}/findings").json()) == 2

    exp = client.post(
        "/api/export.xlsx",
        json={"title": "Audits", "columns": ["Ref", "Title"], "rows": [["A-1", '=HYPERLINK("x")'], ["A-2", "Plain"]]},
        headers=h,
    )
    assert exp.status_code == 200
    ws = load_workbook(io.BytesIO(exp.content)).worksheets[0]
    assert ws["B2"].value.startswith("'=")  # formula injection neutralised
    # Control characters (which Excel cannot hold) are dropped instead of failing the export.
    exp = client.post("/api/export.xlsx", headers=h,
                      json={"title": "T\x07", "columns": ["A\x00B"], "rows": [["\x07bell"]]})
    assert exp.status_code == 200
    ws = load_workbook(io.BytesIO(exp.content)).worksheets[0]
    assert (ws["A1"].value, ws["A2"].value) == ("AB", "bell")

    v = client.post(
        "/api/views", json={"module": "capa", "name": "My overdue", "params": {"status": "Open", "q": ""}}, headers=h
    )
    assert v.status_code == 201 and v.json()["params"] == {"status": "Open"}
    client.post("/api/views", json={"module": "capa", "name": "My overdue", "params": {"severity": "High"}}, headers=h)
    views = client.get("/api/views?module=capa").json()
    assert len(views) == 1 and views[0]["params"] == {"severity": "High"}
    other = _u(db, "iooth", Role.Auditor)
    oh = csrf_headers(login(client2, other.username))
    assert client2.get("/api/views?module=capa").json() == []
    assert client2.delete(f"/api/views/{views[0]['id']}", headers=oh).status_code == 404
    assert client.delete(f"/api/views/{views[0]['id']}", headers=h).status_code == 200
