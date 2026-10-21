"""PROJECTBLUEPRINT reporting: Reports (Views, filters, grouping, export), Dashboards, Presentations, sharing by
access, and AI Reporting workstreams (confirmed sources, stored facts, separate narrative, human review)."""

from __future__ import annotations

import io
import json
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core import permissions as perms
from app.core.enums import Role
from app.main import app
from app.services.ai import providers
from tests.conftest import csrf_headers, login, make_user


def _u(db, prefix, job, role=Role.Auditor):
    user = make_user(db, f"{prefix}_{uuid.uuid4().hex[:6]}", role)
    r = perms.JOB_ROLES_BY_CODE[job]
    user.job_role, user.role, user.review_level = r.code, r.tier, r.review_level
    db.commit()
    return user


def _issue(client, h, tag, rating, title):
    audit = client.post("/api/audits", json={"title": f"Rep {tag} {uuid.uuid4().hex[:3]}", "scope": "x"},
                        headers=h).json()
    f = client.post(f"/api/audits/{audit['id']}/findings", json={
        "title": f"{title} {tag}", "description": "d", "risk_rating": rating, "condition": "c",
        "recommendation": "r"}, headers=h)
    assert f.status_code == 201, f.text
    return audit, f.json()


def test_reports_dashboards_presentations_and_sharing(client, client2, db):
    author = _u(db, "rpaut", "audit_manager")
    other = _u(db, "rpoth", "auditor")
    adm = _u(db, "rpadm", "platform_admin", Role.Admin)
    auditee = _u(db, "rpcli", "audit_client", Role.Auditee)
    h = csrf_headers(login(client, author.username))
    tag = uuid.uuid4().hex[:6]
    _issue(client, h, tag, "High", "Fuel")
    _issue(client, h, tag, "High", "Payroll")
    _issue(client, h, tag, "Low", "Stores")

    views = {v["key"]: v for v in client.get("/api/reporting/views").json()}
    assert {"issues", "actions", "audits", "controls", "risks", "projects"} <= set(views)
    spec = {"view": "issues", "columns": ["issue", "title", "rating", "status"],
            "filters": [{"column": "title", "op": "contains", "value": tag}], "group_by": "rating", "chart": "bar",
            "sort": "rating"}
    bad = client.post("/api/reporting/run", json={**spec, "columns": ["nope"]}, headers=h)
    assert bad.status_code == 422
    run = client.post("/api/reporting/run", json=spec, headers=h).json()
    assert run["total"] == 3 and {g["value"]: g["count"] for g in run["groups"]} == {"High": 2, "Low": 1}
    assert run["filters_text"].startswith("Title contains") and run["period_text"] == "All dates"
    assert run["source"].startswith("Issues") and [c["key"] for c in run["columns"]] == spec["columns"]

    rep = client.post("/api/reporting/items", json={"kind": "report", "name": f"High issues {tag}",
                                                     "definition": spec}, headers=h)
    assert rep.status_code == 201, rep.text
    rep = rep.json()
    assert rep["status"] == "Draft"

    # Drafts and Private items are the owner's; others cannot see or run them.
    oh = csrf_headers(login(client2, other.username))
    assert client2.get(f"/api/reporting/items/{rep['id']}/run").status_code == 404
    client.patch(f"/api/reporting/items/{rep['id']}", json={"status": "Active", "access": "Public",
                                                             "folder": "Audit Reports"}, headers=h)
    shared = client2.get(f"/api/reporting/items/{rep['id']}/run").json()
    assert shared["result"]["total"] == 3 and shared["can_edit"] is False
    assert client2.patch(f"/api/reporting/items/{rep['id']}", json={"name": "Mine now"}, headers=oh).status_code == 403

    assert client.patch(f"/api/reporting/items/{rep['id']}", json={"folder": "Nowhere"}, headers=h).status_code == 422
    assert "Audit Reports" in client.get("/api/reporting/folders").json()

    # Private access: only the chosen users and group members (and the owner) see it.
    ah = csrf_headers(login(client2, adm.username))
    grp = client2.post("/api/groups", json={"name": f"Readers {tag}"}, headers=ah).json()
    assert client.patch(f"/api/reporting/items/{rep['id']}", json={"access": "Private", "group_ids": [999_999]},
                        headers=h).status_code == 422
    client.patch(f"/api/reporting/items/{rep['id']}", json={"access": "Private", "group_ids": [grp["id"]]}, headers=h)
    oh = csrf_headers(login(client2, other.username))
    assert client2.get(f"/api/reporting/items/{rep['id']}").status_code == 404
    ah = csrf_headers(login(client2, adm.username))
    client2.put(f"/api/groups/{grp['id']}/members", json={"user_ids": [other.id]}, headers=ah)
    oh = csrf_headers(login(client2, other.username))
    assert client2.get(f"/api/reporting/items/{rep['id']}").status_code == 200

    # Auditees do not use the reporting app.
    with TestClient(app) as ac:
        login(ac, auditee.username)
        assert ac.get("/api/reporting/views").status_code == 403

    # Export (Excel and CSV), formula-safe, with an About sheet stating source, filters and period.
    x = client.get(f"/api/reporting/items/{rep['id']}/export?format=xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    cells = [str(c.value) for row in wb["Report"].iter_rows() for c in row if c.value]
    titles = [v for v in cells if v.endswith(f" {tag}") and not v.startswith("High issues")]
    assert sorted(v.split()[0] for v in titles) == ["Fuel", "Payroll", "Stores"]
    assert "Source" in [str(c.value) for row in wb["About"].iter_rows() for c in row]
    c = client.get(f"/api/reporting/items/{rep['id']}/export?format=csv")
    assert c.status_code == 200 and "Fuel" in c.content.decode("utf-8-sig")

    # Email: only people it is shared with can be sent it.
    sent = client.post(f"/api/reporting/items/{rep['id']}/email", json={"user_ids": [other.id], "message": "FYI"},
                       headers=h)
    assert sent.status_code == 200 and sent.json()["sent"] == 1
    stranger = _u(db, "rpstr", "auditor")
    assert client.post(f"/api/reporting/items/{rep['id']}/email", json={"user_ids": [stranger.id]},
                       headers=h).status_code == 422

    # Templates: Make a copy gives a Draft report in My Reports.
    templates = {t["key"]: t for t in client.get("/api/reporting/templates").json()}
    copy = client.post("/api/reporting/templates/issues-owner-severity/copy", headers=h).json()
    assert copy["status"] == "Draft" and copy["folder"] == "My Reports" and copy["name"].startswith("Copy of")
    assert "issues-owner-severity" in templates

    # Dashboard: tabs of widgets (report, text, filter); filters apply to every report widget.
    dash = client.post("/api/reporting/items", json={"kind": "dashboard", "name": f"Board {tag}", "definition": {
        "layout": "canvas", "tabs": [
            {"name": "Overview", "widgets": [{"type": "report", "report_id": rep["id"], "visualization": "chart"},
                                             {"type": "text", "text": "Quarter commentary"},
                                             {"type": "filter", "filter": "project"}]},
            {"name": "Empty", "widgets": []}]}}, headers=h)
    assert dash.status_code == 201, dash.text
    did = dash.json()["id"]
    run = client.get(f"/api/reporting/items/{did}/run").json()
    widgets = run["tabs"][0]["widgets"]
    assert widgets[0]["result"]["total"] == 3 and widgets[1]["text"] == "Quarter commentary"
    one = client.get(f"/api/reporting/items/{did}/run?audit=AUD-999999").json()["tabs"][0]["widgets"][0]
    assert one["result"]["total"] == 0 and "Audit is AUD-999999" in one["result"]["filters_text"]
    # Every tab needs content before activation.
    assert client.patch(f"/api/reporting/items/{did}", json={"status": "Active"}, headers=h).status_code == 409
    assert client.post("/api/reporting/items", json={"kind": "dashboard", "name": "Bad", "definition": {
        "tabs": [{"name": "T", "widgets": [{"type": "report", "report_id": 999_999_999}]}]}},
        headers=h).status_code == 404
    assert client.post("/api/reporting/items", json={"kind": "dashboard", "name": "Bad image", "definition": {
        "tabs": [{"name": "T", "widgets": [{"type": "image", "image": "javascript:alert(1)"}]}]}},
        headers=h).status_code == 422

    # Present: ordered slides with a theme.
    pres = client.post("/api/reporting/items", json={"kind": "presentation", "name": f"Committee {tag}", "definition": {
        "theme": "Harbour", "slides": [
            {"name": "Executive summary", "widgets": [{"type": "text", "text": "Key points"}]},
            {"name": "Open findings", "widgets": [{"type": "report", "report_id": rep["id"]}]}]}}, headers=h)
    assert pres.status_code == 201, pres.text
    slides = client.get(f"/api/reporting/items/{pres.json()['id']}/run").json()["slides"]
    assert [s["name"] for s in slides] == ["Executive summary", "Open findings"]
    assert slides[1]["widgets"][0]["result"]["total"] == 3

    # Archived, not erased.
    gone = client.delete(f"/api/reporting/items/{pres.json()['id']}", headers=h).json()
    assert gone["archived"] is True
    assert pres.json()["id"] in {i["id"] for i in client.get("/api/reporting/items?archived=true").json()}


def test_ai_reporting_facts_review_and_publish(client, client2, db, monkeypatch):
    author = _u(db, "aiaut", "auditor")
    reviewer = _u(db, "airev", "audit_manager")
    adm = _u(db, "aiadm", "platform_admin", Role.Admin)
    h = csrf_headers(login(client, author.username))
    tag = uuid.uuid4().hex[:6]
    _issue(client, h, tag, "High", "Fuel")
    _issue(client, h, tag, "Medium", "Payroll")
    source = {"label": "Issues this cycle", "spec": {"view": "issues", "group_by": "rating",
                                                     "filters": [{"column": "title", "op": "contains", "value": tag}]}}
    # New report: a template or a described request gives suggested sources to review and confirm.
    tpl = {t["key"] for t in client.get("/api/ai-reporting/templates").json()}
    assert "committee-summary" in tpl
    sug = client.post("/api/ai-reporting/suggest", json={"template": "committee-summary"}, headers=h).json()
    assert sug["title"] == "Audit Committee Summary" and len(sug["sources"]) == 3
    asked = client.post("/api/ai-reporting/suggest", json={"request": "Summarize open audit findings by severity"},
                        headers=h).json()
    assert asked["sources"][0]["spec"] == {"view": "issues", "group_by": "rating",
                                           "filters": [{"column": "status", "op": "ne", "value": "Closed"}]}
    assert client.post("/api/ai-reporting/suggest", json={"request": "hello there"}, headers=h).status_code == 422
    w = client.post("/api/ai-reporting", json={"title": f"Cycle report {tag}", "purpose": "Committee pack",
                                               "sources": [source], "access": "Public"}, headers=h)
    assert w.status_code == 201, w.text
    w = w.json()
    assert client.post(f"/api/ai-reporting/{w['id']}/submit", headers=h).status_code == 422  # facts first

    # AI is off by default: Generate report stores the facts alone; the narrative cannot be drafted.
    facts = client.post(f"/api/ai-reporting/{w['id']}/generate", headers=h).json()
    assert facts["ai_used"] is False and facts["narrative"] == ""
    src = facts["facts"]["sources"][0]
    assert src["total"] == 2 and src["filters_text"].startswith("Title contains") and src["period_text"] == "All dates"
    assert {g["value"] for g in src["groups"]} == {"High", "Medium"}
    off = client.post(f"/api/ai-reporting/{w['id']}/narrative", headers=h)
    assert off.status_code == 409 and "facts alone" in off.json()["detail"]

    # With a provider: the narrative is drafted from the stored facts only, labelled as generated.
    ah = csrf_headers(login(client2, adm.username))
    calls = []

    def fake(cfg, system, prompt):
        calls.append(prompt)
        return "Two issues were raised this cycle: one High and one Medium."

    monkeypatch.setattr(providers, "complete", fake)
    assert client2.put("/api/config", json={"values": {"ai_enabled": True, "ai_provider": "openai_compatible"}},
                       headers=ah).status_code == 200
    try:
        # Provider unreachable: Generate report still stores the facts and says why there is no narrative.
        def broken(cfg, system, prompt):
            raise providers.AiError("Could not reach the AI provider")

        monkeypatch.setattr(providers, "complete", broken)
        g = client.post(f"/api/ai-reporting/{w['id']}/generate", headers=h)
        assert g.status_code == 200, g.text
        assert g.json()["ai_used"] is False and "reach" in g.json()["ai_error"] and g.json()["facts"] is not None
        monkeypatch.setattr(providers, "complete", fake)
        drafted = client.post(f"/api/ai-reporting/{w['id']}/narrative", headers=h)
        assert drafted.status_code == 200, drafted.text
    finally:
        client2.post("/api/config/ai_enabled/reset", headers=ah)
        client2.post("/api/config/ai_provider/reset", headers=ah)
    d = drafted.json()
    assert d["narrative_origin"] == "ai" and "Generated" in d["narrative_label"]
    assert "ONLY the facts" in calls[0] and f"Cycle report {tag}" in calls[0]
    edited = client.patch(f"/api/ai-reporting/{w['id']}", json={"narrative": d["narrative"] + " Both are open."},
                          headers=h).json()
    assert edited["narrative_origin"] == "edited" and edited["facts"] is not None

    # Not published without human review by someone else.
    assert client.post(f"/api/ai-reporting/{w['id']}/publish", headers=h).status_code == 409
    client.post(f"/api/ai-reporting/{w['id']}/submit", headers=h)
    assert client.post(f"/api/ai-reporting/{w['id']}/review", json={"decision": "approve", "facts_checked": True},
                       headers=h).status_code == 403
    rh = csrf_headers(login(client2, reviewer.username))
    unchecked = client2.post(f"/api/ai-reporting/{w['id']}/review", json={"decision": "approve"}, headers=rh)
    assert unchecked.status_code == 422
    ok = client2.post(f"/api/ai-reporting/{w['id']}/review", json={"decision": "approve", "facts_checked": True,
                                                                   "note": "Agreed to the register"}, headers=rh)
    assert ok.status_code == 200 and ok.json()["status"] == "Approved"
    pub = client.post(f"/api/ai-reporting/{w['id']}/publish", headers=h)
    assert pub.status_code == 200 and pub.json()["status"] == "Published"
    assert client.patch(f"/api/ai-reporting/{w['id']}", json={"narrative": "Changed"}, headers=h).status_code == 409

    # Readers see the facts only if they may see every audit they came from.
    reader = client2.get(f"/api/ai-reporting/{w['id']}").json()
    assert reader["facts"] is not None and reader["status"] == "Published"
    from sqlalchemy import select

    from app.platform.history.models import AuditLog

    db.rollback()
    logged = list(db.execute(select(AuditLog.action_performed)
                             .where(AuditLog.action_performed.like(f"AIREPORT_% id={w['id']}%"))).scalars())
    assert {x.split()[0] for x in logged} >= {"AIREPORT_CREATE", "AIREPORT_FACTS", "AIREPORT_NARRATIVE",
                                              "AIREPORT_SUBMIT", "AIREPORT_REVIEW", "AIREPORT_PUBLISH"}
    assert json.loads(json.dumps(reader["facts"]))["sources"][0]["label"] == "Issues this cycle"


def test_stored_facts_hidden_from_readers_outside_scope(client, client2, db):
    from app.services.reporting.sharing import may_read_facts

    user = _u(db, "aisc", "auditor")
    facts = {"sources": [{"audit_ids": [999_999_998], "perms": []}]}
    from app.core.appconfig import cfg

    if cfg("port_scope_enforced"):
        assert may_read_facts(db, user, facts) is False
    assert may_read_facts(db, user, {"sources": [{"audit_ids": [], "perms": ["grc.risk"]}]}) is False
