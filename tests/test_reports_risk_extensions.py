"""Reports: Cross Tab, Sections, Analysis Style and Design; broadcast schedules; slide videos; presentation export
to PowerPoint and PDF. Risk: assessment driven by key risk indicators, and per-segment assessments with roll-up."""

from __future__ import annotations

import io
import uuid
import zipfile
from datetime import timedelta

from app.core import permissions as perms
from app.core.enums import Role, utcnow
from app.platform.notifications.models import Notification
from app.services.reporting.broadcasts import run_broadcasts
from app.services.reporting.models import ReportBroadcast
from tests.conftest import csrf_headers, login, make_user


def _u(db, prefix, job, role=Role.Auditor):
    user = make_user(db, f"{prefix}_{uuid.uuid4().hex[:6]}", role)
    r = perms.JOB_ROLES_BY_CODE[job]
    user.job_role, user.role, user.review_level = r.code, r.tier, r.review_level
    db.commit()
    return user


def _issues(client, h, tag):
    for rating, title in (("High", "Fuel"), ("High", "Payroll"), ("Low", "Stores")):
        audit = client.post("/api/audits", json={"title": f"Ext {tag} {title}", "scope": "x"}, headers=h).json()
        f = client.post(f"/api/audits/{audit['id']}/findings", json={
            "title": f"{title} {tag}", "description": "d", "risk_rating": rating, "condition": "c",
            "recommendation": "r"}, headers=h)
        assert f.status_code == 201, f.text


def test_crosstab_sections_design_and_drill(client, db):
    author = _u(db, "xtaut", "audit_manager")
    h = csrf_headers(login(client, author.username))
    tag = uuid.uuid4().hex[:6]
    _issues(client, h, tag)
    base = {"view": "issues", "columns": ["issue", "title", "rating", "status"],
            "filters": [{"column": "title", "op": "contains", "value": tag}]}

    ct = client.post("/api/reporting/run", json={**base, "layout": "crosstab", "group_by": "rating",
                                                 "cross_column": "status", "chart": "bar"}, headers=h)
    assert ct.status_code == 200, ct.text
    ct = ct.json()
    assert ct["chart"] is None and ct["layout"] == "crosstab"
    x = ct["crosstab"]
    assert x["measure"] == "count" and x["grand"] == 3
    assert {r["value"]: r["total"] for r in x["rows"]} == {"High": 2, "Low": 1}
    assert sum(x["totals"]) == 3
    # A Cross Tab needs two different fields; the value must be a number column.
    assert client.post("/api/reporting/run", json={**base, "layout": "crosstab", "group_by": "rating"},
                       headers=h).status_code == 422
    assert client.post("/api/reporting/run", json={**base, "layout": "crosstab", "group_by": "rating",
                                                   "cross_column": "rating"}, headers=h).status_code == 422
    assert client.post("/api/reporting/run", json={**base, "layout": "crosstab", "group_by": "rating",
                                                   "cross_column": "status", "measure": "title"},
                       headers=h).status_code == 422

    sec = client.post("/api/reporting/run", json={**base, "section_by": "rating", "analysis": "drill",
                                                  "design": {"font": "Serif", "borders": "Grid"}}, headers=h).json()
    assert {s["value"]: s["count"] for s in sec["sections"]} == {"High": 2, "Low": 1}
    assert sum(len(s["rows"]) for s in sec["sections"]) == 3
    assert sec["analysis"] == "drill" and sec["design"] == {"font": "Serif", "size": "Medium", "borders": "Grid",
                                                             "align": "Left", "page_layout": "Portrait"}
    bad = client.post("/api/reporting/run", json={**base, "design": {"font": "Comic"}}, headers=h)
    assert bad.status_code == 422

    # Saved with the layout; the Excel export carries a Cross Tab sheet.
    it = client.post("/api/reporting/items", json={"kind": "report", "name": f"Matrix {tag}", "definition": {
        **base, "layout": "crosstab", "group_by": "rating", "cross_column": "status"}}, headers=h).json()
    assert it["definition"]["layout"] == "crosstab"
    xl = client.get(f"/api/reporting/items/{it['id']}/export?format=xlsx")
    assert xl.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames[0] == "Cross Tab"
    cells = [str(c.value) for row in wb["Cross Tab"].iter_rows() for c in row if c.value is not None]
    assert any(" / Status" in c for c in cells) and "Total" in cells


def test_broadcast_schedule_sends_to_shared_people_only(client, db):
    author = _u(db, "bcaut", "audit_manager")
    shared = _u(db, "bcshr", "auditor")
    outsider = _u(db, "bcout", "auditor")
    h = csrf_headers(login(client, author.username))
    it = client.post("/api/reporting/items", json={
        "kind": "report", "name": f"Weekly issues {uuid.uuid4().hex[:5]}", "definition": {"view": "issues"},
        "access": "Private", "user_ids": [shared.id]}, headers=h).json()
    body = {"frequency": "Weekly", "weekday": 0, "hour": 7, "user_ids": [shared.id, outsider.id],
            "message": "For the Monday meeting"}
    assert client.post(f"/api/reporting/items/{it['id']}/broadcasts", json=body, headers=h).status_code == 409
    client.patch(f"/api/reporting/items/{it['id']}", json={"status": "Active"}, headers=h)
    missing_day = client.post(f"/api/reporting/items/{it['id']}/broadcasts", json={**body, "weekday": None},
                              headers=h)
    assert missing_day.status_code == 422
    b = client.post(f"/api/reporting/items/{it['id']}/broadcasts", json=body, headers=h)
    assert b.status_code == 201, b.text
    b = b.json()
    assert b["next_run_at"] is not None and b["active"] is True
    assert [x["id"] for x in client.get(f"/api/reporting/items/{it['id']}/broadcasts").json()] == [b["id"]]

    # The timer sends it: the shared user is notified; the outsider is skipped.
    row = db.get(ReportBroadcast, b["id"])
    row.next_run_at = utcnow() - timedelta(minutes=5)
    db.commit()
    assert run_broadcasts(db) >= 1
    db.expire_all()
    got = {n.user_id for n in db.query(Notification).filter(Notification.link_type == "report_item",
                                                             Notification.link_id == it["id"])}
    assert shared.id in got and outsider.id not in got
    row = db.get(ReportBroadcast, b["id"])
    assert row.last_result.startswith("Sent to 1") and "1 not shared" in row.last_result
    assert row.next_run_at > utcnow()

    # Archived: the schedule pauses itself.
    client.delete(f"/api/reporting/items/{it['id']}", headers=h)
    row.next_run_at = utcnow() - timedelta(minutes=1)
    db.commit()
    run_broadcasts(db)
    db.expire_all()
    row = db.get(ReportBroadcast, b["id"])
    assert row.active is False and row.last_result.startswith("Paused")


def test_video_slides_and_presentation_export(client, client2, db):
    author = _u(db, "pxaut", "audit_manager")
    viewer = _u(db, "pxvw", "auditor")
    h = csrf_headers(login(client, author.username))
    tag = uuid.uuid4().hex[:6]
    _issues(client, h, tag)
    rep = client.post("/api/reporting/items", json={
        "kind": "report", "name": f"By severity {tag}", "access": "Public", "definition": {
            "view": "issues", "filters": [{"column": "title", "op": "contains", "value": tag}],
            "group_by": "rating", "chart": "bar"}}, headers=h).json()
    client.patch(f"/api/reporting/items/{rep['id']}", json={"status": "Active"}, headers=h)

    fake = b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64
    up = client.post("/api/reporting/media", files={"file": ("clip.mp4", fake, "video/mp4")}, headers=h)
    assert up.status_code == 201, up.text
    vid = up.json()
    renamed = client.post("/api/reporting/media", files={"file": ("clip.mp4", b"not a video at all", "video/mp4")},
                          headers=h)
    assert renamed.status_code == 415
    assert client.post("/api/reporting/media", files={"file": ("clip.avi", fake, "video/x-msvideo")},
                       headers=h).status_code == 415

    png = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8/5+hHgAHggJ/PchI7wAAAA"
           "BJRU5ErkJggg==")
    slides = [{"name": "Executive summary", "widgets": [
        {"type": "text", "text": "Key themes … and decisions ≥ last quarter", "x": 0, "y": 0, "w": 6, "h": 3},
        {"type": "report", "report_id": rep["id"], "visualization": "chart", "x": 6, "y": 0, "w": 6, "h": 4},
        {"type": "image", "image": png, "x": 0, "y": 3, "w": 3, "h": 2}]},
        {"name": "Walkthrough video", "widgets": [{"type": "video", "video_id": vid["id"], "x": 0, "y": 0, "w": 8,
                                                   "h": 5}]},
        {"name": "Table", "widgets": [{"type": "report", "report_id": rep["id"], "visualization": "table",
                                       "x": 0, "y": 0, "w": 12, "h": 6}]}]
    pres = client.post("/api/reporting/items", json={"kind": "presentation", "name": f"Quarterly {tag}",
                                                     "access": "Public",
                                                     "definition": {"theme": "Harbour", "slides": slides}}, headers=h)
    assert pres.status_code == 201, pres.text
    pres = pres.json()
    bad_video = client.post("/api/reporting/items", json={"kind": "presentation", "name": "Bad video", "definition": {
        "slides": [{"name": "x", "widgets": [{"type": "video", "video_id": 999999}]}]}}, headers=h)
    assert bad_video.status_code == 422

    pptx = client.get(f"/api/reporting/items/{pres['id']}/export/pptx")
    assert pptx.status_code == 200, pptx.text
    assert pptx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    names = zipfile.ZipFile(io.BytesIO(pptx.content)).namelist()
    assert sum(1 for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")) == 3
    assert any(n.startswith("ppt/charts/") for n in names) and any(n.startswith("ppt/media/") for n in names)
    pdf = client.get(f"/api/reporting/items/{pres['id']}/export/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == \
        "application/pdf"
    assert client.get(f"/api/reporting/items/{rep['id']}/export/pdf").status_code == 422

    # The viewer cannot fetch the video until the presentation is active and shared with them.
    vh = csrf_headers(login(client2, viewer.username))
    assert client2.get(f"/api/reporting/media/{vid['id']}").status_code == 404
    client.patch(f"/api/reporting/items/{pres['id']}", json={"status": "Active"}, headers=h)
    got = client2.get(f"/api/reporting/media/{vid['id']}")
    assert got.status_code == 200 and got.content == fake and got.headers["content-type"] == "video/mp4"
    assert client2.get(f"/api/reporting/items/{pres['id']}/export/pptx").status_code == 200
    assert vh


def test_kri_driven_assessment_and_segments(client, client2, db):
    mgr = _u(db, "krmgr", "audit_manager")
    adm = _u(db, "kradm", "platform_admin", Role.Admin)
    h = csrf_headers(login(client, mgr.username))
    ah = csrf_headers(login(client2, adm.username))
    tag = uuid.uuid4().hex[:5]
    r = client.post("/api/grc/risks", json={"title": f"Berth congestion {tag}", "category": "Operational",
                                             "inherent_likelihood": 5, "inherent_impact": 4,
                                             "residual_likelihood": 2, "residual_impact": 3}, headers=h)
    assert r.status_code == 201, r.text
    risk = r.json()
    k = client.post(f"/api/grc/risks/{risk['id']}/kris", json={"name": f"Vessel wait hours {tag}", "unit": "h",
                                                                "amber_at": 24, "red_at": 48}, headers=h).json()
    d = client.put(f"/api/grc/kris/{k['id']}/driver", json={"drives": "likelihood", "amber_step": 1, "red_step": 3},
                   headers=h)
    assert d.status_code == 200, d.text
    assert client.put(f"/api/grc/kris/{k['id']}/driver", json={"drives": "impact", "amber_step": 2, "red_step": 1},
                      headers=h).status_code == 422

    none = client.get(f"/api/grc/risks/{risk['id']}/kri-assessment").json()
    assert none["changes"] is False and none["drivers"][0]["rag"] == "No data"
    assert client.post(f"/api/grc/risks/{risk['id']}/kri-assessment", headers=h).status_code == 409

    today = utcnow().date().isoformat()
    client.post(f"/api/grc/kris/{k['id']}/readings", json={"period_date": today, "value": 60}, headers=h)
    p = client.get(f"/api/grc/risks/{risk['id']}/kri-assessment").json()
    assert p["drivers"][0]["rag"] == "Red" and p["drivers"][0]["step"] == 3
    assert p["proposed"]["residual_likelihood"] == 5 and p["proposed"]["residual_impact"] == 3  # 2 + 3, cap 5
    applied = client.post(f"/api/grc/risks/{risk['id']}/kri-assessment", headers=h)
    assert applied.status_code == 200, applied.text
    full = client.get(f"/api/grc/risks/{risk['id']}").json()
    assert full["residual_likelihood"] == 5 and full["assessments"][0]["origin"] == "kri"
    assert full["kris"][0]["drives"] == "likelihood"

    # Back to Green: the proposal returns to the last assessment a person made.
    client.post(f"/api/grc/kris/{k['id']}/readings", json={"period_date": today, "value": 5}, headers=h)
    back = client.get(f"/api/grc/risks/{risk['id']}/kri-assessment").json()
    assert back["proposed"]["residual_likelihood"] == 2 and back["changes"] is True

    # Automatic when the administrator turns it on.
    assert client2.put("/api/config", json={"values": {"erm_kri_auto_assess": True}}, headers=ah).status_code == 200
    try:
        client.post(f"/api/grc/kris/{k['id']}/readings", json={"period_date": today, "value": 30}, headers=h)
        auto = client.get(f"/api/grc/risks/{risk['id']}").json()
        assert auto["residual_likelihood"] == 3 and auto["assessments"][0]["origin"] == "kri"
    finally:
        client2.post("/api/config/erm_kri_auto_assess/reset", headers=ah)

    # Segments: the latest assessment of each segment rolls up (worst by default, or average).
    orgs = []
    for name in ("Tema", "Takoradi"):
        o = client2.post("/api/organizations", json={"name": f"{name} {tag}", "category": "Port"}, headers=ah)
        assert o.status_code in (200, 201), o.text
        orgs.append(o.json()["id"])
    assert client.post(f"/api/grc/risks/{risk['id']}/segments/rollup", headers=h).status_code == 422
    s1 = client.post(f"/api/grc/risks/{risk['id']}/segments", json={
        "org_id": orgs[0], "inherent_likelihood": 5, "inherent_impact": 4, "residual_likelihood": 4,
        "residual_impact": 4}, headers=h)
    assert s1.status_code == 201, s1.text
    client.post(f"/api/grc/risks/{risk['id']}/segments", json={
        "org_id": orgs[1], "inherent_likelihood": 3, "inherent_impact": 3, "residual_likelihood": 2,
        "residual_impact": 2}, headers=h)
    bad = client.post(f"/api/grc/risks/{risk['id']}/segments", json={
        "org_id": orgs[1], "inherent_likelihood": 2, "inherent_impact": 2, "residual_likelihood": 3,
        "residual_impact": 3}, headers=h)
    assert bad.status_code == 422
    seg = client.get(f"/api/grc/risks/{risk['id']}/segments").json()
    assert seg["method"] == "Worst" and len(seg["segments"]) == 2
    assert seg["rollup"]["residual_score"] == 16 and seg["rollup"]["inherent_score"] == 20
    done = client.post(f"/api/grc/risks/{risk['id']}/segments/rollup", headers=h)
    assert done.status_code == 200
    full = client.get(f"/api/grc/risks/{risk['id']}").json()
    assert full["residual_score"] == 16 and full["assessments"][0]["origin"] == "segments"
    assert client2.put("/api/config", json={"values": {"erm_segment_rollup": "Average"}}, headers=ah).status_code == 200
    try:
        avg = client.get(f"/api/grc/risks/{risk['id']}/segments").json()["rollup"]
        assert (avg["inherent_likelihood"], avg["inherent_impact"], avg["residual_likelihood"],
                avg["residual_impact"]) == (4, 4, 3, 3)
    finally:
        client2.post("/api/config/erm_segment_rollup/reset", headers=ah)


def test_risk_state_move_on_profile(client, db):
    mgr = _u(db, "stmv", "audit_manager")
    aud = _u(db, "staud2", "auditor")
    h = csrf_headers(login(client, mgr.username))
    r = client.post("/api/grc/risks", json={"title": f"State move {uuid.uuid4().hex[:5]}", "inherent_likelihood": 3,
                                             "inherent_impact": 3, "residual_likelihood": 2, "residual_impact": 2},
                    headers=h).json()
    moved = client.patch(f"/api/grc/risks/{r['id']}/state", json={"treatment": "mitigate"}, headers=h)
    assert moved.status_code == 200 and moved.json()["treatment"] == "Mitigate"
    assert client.patch(f"/api/grc/risks/{r['id']}/state", json={"treatment": "Ignore"}, headers=h).status_code == 422
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as other:
        oh = csrf_headers(login(other, aud.username))
        assert other.patch(f"/api/grc/risks/{r['id']}/state", json={"treatment": ""}, headers=oh).status_code == 403
