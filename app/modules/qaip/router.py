"""QAIP: per-engagement quality review checklist and post-audit satisfaction survey."""

from __future__ import annotations

import json
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Request, status
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession

from app.core.access import ensure_can_view
from app.core.appconfig import cfg
from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import require_active_user, require_module_tier, staff_or_perm
from app.core.enums import AuditStatus, ReportStatus, Role, TestStatus, utcnow
from app.core.errors import field_error
from app.core.history import log_record_history
from app.modules.audits.models import Audit, AuditAuditee, Finding
from app.modules.qaip.models import AuditSurvey, QualityReview
from app.modules.users.models import User

router = APIRouter(tags=["qaip"])
Staff = Depends(require_module_tier("qaip", Role.Admin, Role.Auditor))
Governance = Depends(staff_or_perm("governance.view"))

# (key, standard, question)
CHECKLIST: list[tuple[str, str, str]] = [
    ("objectives", "GIAS 13.3", "Engagement objectives, scope and criteria were documented"),
    ("risk_assessment", "GIAS 13.2", "A risk assessment of the area informed the work programme"),
    ("memo_approved", "GIAS 13.6", "The planning memo / work programme was approved before fieldwork"),
    ("independence", "GIAS 2.1", "Independence and objectivity of the team were considered and documented"),
    ("evidence", "GIAS 14.1", "Working papers hold sufficient, relevant and reliable evidence"),
    ("supervision", "GIAS 12.3", "Test work was reviewed and signed off by someone other than the performer"),
    ("finding_attributes", "GIAS 14.3", "Each finding states criteria, condition, cause, effect and recommendation"),
    ("management_discussion", "GIAS 14.4", "Findings were discussed with management and responses obtained"),
    ("report_issued", "GIAS 15.1", "The final report was approved by the CAE and issued"),
    ("action_plans", "GIAS 15.2", "Agreed actions have owners and due dates"),
    ("budget", "GIAS 13.5", "The engagement was completed within (or explained against) its budget"),
    ("retention", "GIAS 14.6", "Engagement records are retained with controlled access"),
]
ANSWERS = ("Conforms", "Partially conforms", "Does not conform", "N/A")

SURVEY: list[tuple[str, str]] = [
    ("communication", "The objectives and scope of the audit were explained clearly"),
    ("professionalism", "The audit team was professional and courteous"),
    ("disruption", "The audit was carried out with minimal disruption to operations"),
    ("accuracy", "Findings were accurate, clear and discussed before the report"),
    ("usefulness", "Recommendations were practical and add value"),
    ("timeliness", "The report was issued in good time"),
]


def _audit(db: DBSession, audit_id: int, user: User) -> Audit:
    audit = db.get(Audit, audit_id)
    ensure_can_view(db, user, "audit", audit)
    return audit


# ------------------------------------------------------------------ quality review


def evidence_hints(db: DBSession, audit: Audit) -> dict[str, str]:
    """What the system already knows for each checklist line (the reviewer still decides)."""
    from app.modules.capa.models import Capa
    from app.modules.capacity.router import audit_budgets
    from app.modules.memos.models import PlanningMemo
    from app.modules.rcm.models import Control, ControlTest, Objective
    from app.modules.reports.models import AuditReport

    hints: dict[str, str] = {}
    memo = db.execute(select(PlanningMemo).where(PlanningMemo.audit_id == audit.id)).scalar_one_or_none()
    objectives = db.execute(select(Objective).where(Objective.audit_id == audit.id)).scalars().all()
    hints["objectives"] = f"{len(objectives)} objective(s) recorded" + (
        "; memo has objectives" if memo and memo.objectives.strip() else ""
    )
    hints["risk_assessment"] = (
        "Key risks documented in the planning memo"
        if memo and memo.key_risks.strip()
        else "No key risks in the planning memo"
    )
    hints["memo_approved"] = f"Planning memo {memo.status.lower()}" if memo else "No planning memo"
    tests = (
        db.execute(
            select(ControlTest)
            .join(Control, Control.id == ControlTest.control_id)
            .join(Objective, Objective.id == Control.objective_id)
            .where(Objective.audit_id == audit.id)
        )
        .scalars()
        .all()
    )
    reviewed = [t for t in tests if t.status == TestStatus.Reviewed]
    hints["supervision"] = f"{len(reviewed)} of {len(tests)} test(s) reviewed" if tests else "No control tests"
    findings = db.execute(select(Finding).where(Finding.audit_id == audit.id)).scalars().all()
    complete = [
        f
        for f in findings
        if all((getattr(f, k) or "").strip() for k in ("criteria", "condition", "cause", "effect", "recommendation"))
    ]
    hints["finding_attributes"] = (
        f"{len(complete)} of {len(findings)} finding(s) fully documented" if findings else "No findings"
    )
    responded = [f for f in findings if (f.management_response or "").strip()]
    hints["management_discussion"] = (
        f"{len(responded)} of {len(findings)} with a management response" if findings else "No findings"
    )
    report = db.execute(select(AuditReport).where(AuditReport.audit_id == audit.id)).scalar_one_or_none()
    hints["report_issued"] = (
        "Report issued"
        if report and report.status == ReportStatus.Issued
        else (f"Report in {report.status.value}" if report else "No report")
    )
    actions = (
        db.execute(
            select(Capa).where(Capa.source_type == "finding", Capa.source_id.in_([f.id for f in findings] or [0]))
        )
        .scalars()
        .all()
    )
    owned = [c for c in actions if c.owner_email and c.due_date]
    hints["action_plans"] = (
        f"{len(owned)} of {len(actions)} action(s) with owner and due date" if actions else "No actions raised"
    )
    b = next((x for x in audit_budgets(db) if x.audit_id == audit.id), None)
    hints["budget"] = (
        f"{b.actual_hours:g}h charged of {b.budget_hours}h budget" if b and b.budget_hours else "No budget recorded"
    )
    return hints


class Answer(BaseModel):
    answer: str = ""
    comment: str = Field("", max_length=2000)

    @field_validator("answer")
    @classmethod
    def _a(cls, v: str) -> str:
        if v and v not in ANSWERS:
            raise ValueError(f"must be one of: {', '.join(ANSWERS)}")
        return v


class ReviewIn(BaseModel):
    answers: dict[str, Answer] = Field(default_factory=dict)
    conclusion: str = Field("", max_length=5000)
    complete: bool = False


class ChecklistItem(BaseModel):
    key: str
    standard: str
    question: str
    answer: str = ""
    comment: str = ""
    hint: str = ""


class ReviewOut(BaseModel):
    audit_id: int
    items: list[ChecklistItem]
    answer_options: list[str] = list(ANSWERS)
    conclusion: str = ""
    score: float | None = None
    reviewer_name: str | None = None
    completed_at: datetime | None = None
    can_edit: bool = False


def score_of(answers: dict[str, dict]) -> float | None:
    weights = {"Conforms": 1.0, "Partially conforms": 0.5, "Does not conform": 0.0}
    applicable = [weights[a["answer"]] for a in answers.values() if a.get("answer") in weights]
    return round(100 * sum(applicable) / len(applicable), 1) if applicable else None


def _review_out(db: DBSession, audit: Audit, review: QualityReview | None, user: User) -> ReviewOut:
    answers = json.loads(review.answers) if review else {}
    hints = evidence_hints(db, audit)
    items = [
        ChecklistItem(
            key=k,
            standard=s,
            question=q,
            answer=answers.get(k, {}).get("answer", ""),
            comment=answers.get(k, {}).get("comment", ""),
            hint=hints.get(k, ""),
        )
        for k, s, q in CHECKLIST
    ]
    reviewer = db.get(User, review.reviewer_id) if review and review.reviewer_id else None
    staff = user.role in (Role.Admin, Role.Auditor)
    return ReviewOut(
        audit_id=audit.id,
        items=items,
        conclusion=review.conclusion if review else "",
        score=review.score if review else None,
        reviewer_name=(reviewer.full_name or reviewer.username) if reviewer else None,
        completed_at=review.completed_at if review else None,
        can_edit=staff and audit.created_by != user.id and not (review and review.completed_at),
    )


@router.get("/api/audits/{audit_id}/quality-review", response_model=ReviewOut)
def get_review(audit_id: int, db: DBSession = Depends(get_db), user: User = Staff) -> ReviewOut:
    audit = _audit(db, audit_id, user)
    review = db.execute(select(QualityReview).where(QualityReview.audit_id == audit_id)).scalar_one_or_none()
    return _review_out(db, audit, review, user)


@router.put("/api/audits/{audit_id}/quality-review", response_model=ReviewOut)
def save_review(
    audit_id: int, payload: ReviewIn, request: Request, db: DBSession = Depends(get_db), user: User = Staff
) -> ReviewOut:
    audit = _audit(db, audit_id, user)
    if audit.created_by == user.id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="The quality review must be done by someone who did not prepare the audit",
        )
    if audit.status not in (AuditStatus.Review, AuditStatus.Closed):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="The quality review is done once the audit is in Review or Closed",
        )
    review = db.execute(select(QualityReview).where(QualityReview.audit_id == audit_id)).scalar_one_or_none()
    if review and review.completed_at:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="The quality review is already completed")
    keys = {k for k, _, _ in CHECKLIST}
    unknown = set(payload.answers) - keys
    if unknown:
        raise field_error("answers", f"unknown checklist item(s): {', '.join(sorted(unknown))}")
    answers = {k: v.model_dump() for k, v in payload.answers.items()}
    if payload.complete:
        blank = [q for k, _, q in CHECKLIST if not answers.get(k, {}).get("answer")]
        if blank:
            raise field_error("answers", f"answer every item before completing ({len(blank)} left)")
        gaps = [
            k
            for k, a in answers.items()
            if a["answer"] in ("Partially conforms", "Does not conform") and not a["comment"].strip()
        ]
        if gaps:
            raise field_error("answers", "explain every 'Partially' or 'Does not conform' answer in its comment")
        if not payload.conclusion.strip():
            raise field_error("conclusion", "is required to complete the review")
    if review is None:
        review = QualityReview(audit_id=audit_id)
        db.add(review)
    review.answers = json.dumps(answers)
    review.conclusion = payload.conclusion
    review.score = score_of(answers)
    review.reviewer_id = user.id
    if payload.complete:
        review.completed_at = utcnow()
        log_record_history(
            db,
            record_type="audit",
            record_id=audit.id,
            user_id=user.id,
            action=f"QUALITY REVIEW COMPLETED score={review.score}",
            request=request,
        )
    log_action(db, user_id=user.id, action=f"QA_REVIEW_SAVE audit={audit_id}", request=request)
    db.commit()
    return _review_out(db, audit, review, user)


# ------------------------------------------------------------------ satisfaction survey


class SurveyIn(BaseModel):
    scores: dict[str, int]
    comments: str = Field("", max_length=5000)

    @field_validator("scores")
    @classmethod
    def _scores(cls, v: dict[str, int]) -> dict[str, int]:
        keys = {k for k, _ in SURVEY}
        if set(v) != keys:
            raise ValueError("rate every question")
        if any(not 1 <= s <= 5 for s in v.values()):
            raise ValueError("each rating must be from 1 to 5")
        return v


class SurveyState(BaseModel):
    audit_id: int
    questions: list[dict[str, str]]
    open: bool
    can_respond: bool
    my_response: dict | None = None
    responses: int = 0
    invited: int = 0
    average: float | None = None
    by_question: dict[str, float] = {}
    comments: list[str] = []


def survey_open(db: DBSession, audit: Audit) -> bool:
    from app.modules.reports.models import AuditReport

    if audit.status == AuditStatus.Closed:
        return True
    report = db.execute(select(AuditReport).where(AuditReport.audit_id == audit.id)).scalar_one_or_none()
    return report is not None and report.status == ReportStatus.Issued


def _invited(db: DBSession, audit_id: int) -> list[int]:
    return list(db.execute(select(AuditAuditee.user_id).where(AuditAuditee.audit_id == audit_id)).scalars())


@router.get("/api/audits/{audit_id}/survey", response_model=SurveyState)
def get_survey(
    audit_id: int, db: DBSession = Depends(get_db), user: User = Depends(require_active_user)
) -> SurveyState:
    audit = _audit(db, audit_id, user)
    rows = db.execute(select(AuditSurvey).where(AuditSurvey.audit_id == audit_id)).scalars().all()
    mine = next((r for r in rows if r.respondent_id == user.id), None)
    invited = _invited(db, audit_id)
    state = SurveyState(
        audit_id=audit_id,
        questions=[{"key": k, "text": t} for k, t in SURVEY],
        open=survey_open(db, audit),
        can_respond=survey_open(db, audit) and user.id in invited and mine is None,
        my_response={"scores": json.loads(mine.scores), "comments": mine.comments} if mine else None,
        invited=len(invited),
    )
    if user.role in (Role.Admin, Role.Auditor):
        state.responses = len(rows)
        if rows:
            state.average = round(sum(r.average for r in rows) / len(rows), 2)
            per = [json.loads(r.scores) for r in rows]
            state.by_question = {k: round(sum(p[k] for p in per) / len(per), 2) for k, _ in SURVEY}
            state.comments = [r.comments for r in rows if r.comments.strip()]
    return state


@router.post("/api/audits/{audit_id}/survey", response_model=SurveyState, status_code=status.HTTP_201_CREATED)
def submit_survey(
    audit_id: int,
    payload: SurveyIn,
    request: Request,
    db: DBSession = Depends(get_db),
    user: User = Depends(require_active_user),
) -> SurveyState:
    audit = _audit(db, audit_id, user)
    if user.id not in _invited(db, audit_id):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Only the audit's auditees answer the survey")
    if not survey_open(db, audit):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="The survey opens when the report is issued or the audit is closed",
        )
    if db.execute(
        select(AuditSurvey).where(AuditSurvey.audit_id == audit_id, AuditSurvey.respondent_id == user.id)
    ).first():
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="You have already answered this survey")
    avg = round(sum(payload.scores.values()) / len(payload.scores), 2)
    db.add(
        AuditSurvey(
            audit_id=audit_id,
            respondent_id=user.id,
            scores=json.dumps(payload.scores),
            average=avg,
            comments=payload.comments.strip(),
        )
    )
    log_action(db, user_id=user.id, action=f"SURVEY_SUBMIT audit={audit_id} avg={avg}", request=request)
    db.commit()
    return get_survey(audit_id, db, user)


def invite_survey(db: DBSession, audit: Audit) -> int:
    """Called when an audit closes: ask its auditees for feedback."""
    if not cfg("survey_on_close"):
        return 0
    from app.modules.notifications.service import notify_many

    return notify_many(
        db,
        _invited(db, audit.id),
        f"Tell us how the audit went: {audit.title}",
        "Please take two minutes to rate the audit. Open the audit and choose the Survey tab.",
        link_type="audit",
        link_id=audit.id,
    )


# ------------------------------------------------------------------ QAIP summary


class QaipSummary(BaseModel):
    fiscal_year: int
    closed_audits: int
    reviews_completed: int
    review_coverage_pct: int
    average_score: float | None
    weakest_items: list[dict]
    survey_responses: int
    survey_invited: int
    survey_average: float | None
    survey_by_question: list[dict]
    recent_comments: list[str]
    reviews: list[dict]


def qaip_summary(db: DBSession, fy: int) -> QaipSummary:
    start, end = datetime(fy, 1, 1), datetime(fy + 1, 1, 1)
    audits = db.execute(select(Audit).where(Audit.updated_at >= start, Audit.updated_at < end)).scalars().all()
    ids = [a.id for a in audits]
    closed = [a for a in audits if a.status == AuditStatus.Closed]
    reviews = (
        db.execute(
            select(QualityReview).where(QualityReview.audit_id.in_(ids or [0]), QualityReview.completed_at.is_not(None))
        )
        .scalars()
        .all()
    )
    scores = [r.score for r in reviews if r.score is not None]
    per_item: dict[str, list[float]] = {}
    weights = {"Conforms": 1.0, "Partially conforms": 0.5, "Does not conform": 0.0}
    for r in reviews:
        for k, a in json.loads(r.answers).items():
            if a.get("answer") in weights:
                per_item.setdefault(k, []).append(weights[a["answer"]])
    labels = {k: (s, q) for k, s, q in CHECKLIST}
    weakest = sorted(
        (
            {
                "key": k,
                "standard": labels[k][0],
                "question": labels[k][1],
                "conformance_pct": round(100 * sum(v) / len(v)),
            }
            for k, v in per_item.items()
            if k in labels
        ),
        key=lambda x: x["conformance_pct"],
    )[:5]
    surveys = db.execute(select(AuditSurvey).where(AuditSurvey.audit_id.in_(ids or [0]))).scalars().all()
    invited = sum(len(_invited(db, a.id)) for a in audits if survey_open(db, a))
    per_q = [json.loads(s.scores) for s in surveys]
    by_q = (
        [{"key": k, "question": t, "average": round(sum(p[k] for p in per_q) / len(per_q), 2)} for k, t in SURVEY]
        if per_q
        else []
    )
    titles = {a.id: a.title for a in audits}
    return QaipSummary(
        fiscal_year=fy,
        closed_audits=len(closed),
        reviews_completed=len(reviews),
        review_coverage_pct=round(100 * len([r for r in reviews if r.audit_id in {a.id for a in closed}]) / len(closed))
        if closed
        else 0,
        average_score=round(sum(scores) / len(scores), 1) if scores else None,
        weakest_items=weakest,
        survey_responses=len(surveys),
        survey_invited=invited,
        survey_average=round(sum(s.average for s in surveys) / len(surveys), 2) if surveys else None,
        survey_by_question=by_q,
        recent_comments=[
            s.comments for s in sorted(surveys, key=lambda s: s.submitted_at, reverse=True) if s.comments.strip()
        ][:8],
        reviews=[
            {
                "audit_id": r.audit_id,
                "title": titles.get(r.audit_id, ""),
                "score": r.score,
                "completed_at": r.completed_at.isoformat() if r.completed_at else None,
            }
            for r in reviews
        ],
    )


@router.get("/api/qaip/summary", response_model=QaipSummary)
def get_qaip_summary(
    fiscal_year: int | None = None, db: DBSession = Depends(get_db), user: User = Governance
) -> QaipSummary:
    return qaip_summary(db, fiscal_year or utcnow().year)


# ------------------------------------------------------------------ internal assessment report (GIAS 8.3)


def conformance_rating(score: float | None) -> str:
    """IIA rating scale used for quality assessments."""
    if score is None:
        return "Not assessed"
    if score >= 85:
        return "Generally conforms"
    if score >= 70:
        return "Partially conforms"
    return "Does not conform"


@router.get("/api/qaip/report.docx")
def internal_assessment_report(request: Request, fiscal_year: int | None = None, db: DBSession = Depends(get_db),
                               user: User = Governance):  # noqa: ANN201
    """Periodic internal assessment report for the CAE and the Audit Committee (Word)."""
    import io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from fastapi.responses import Response

    from app.modules.insights.router import analytics

    fy = fiscal_year or utcnow().year
    q = qaip_summary(db, fy)
    a = analytics(db, fy)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.add_run(f"{cfg('org_name').upper()}\n{cfg('department_name').upper()}").bold = True
    doc.add_heading(f"Internal Quality Assessment Report - FY{fy}", level=0)
    doc.add_paragraph(f"Prepared {utcnow():%d %B %Y} under the Quality Assurance and Improvement Programme "
                      "(IIA Global Internal Audit Standards, Principle 8, Standard 8.3).")

    def table(headers: list[str], rows: list[list]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for cell, text in zip(t.rows[0].cells, headers, strict=True):
            cell.text = text
        for row in rows:
            for cell, text in zip(t.add_row().cells, row, strict=True):
                cell.text = str(text)

    rating = conformance_rating(q.average_score)
    doc.add_heading("1. Overall conclusion", level=1)
    doc.add_paragraph(
        f"Based on {q.reviews_completed} engagement quality review(s) covering {q.review_coverage_pct}% of the "
        f"{q.closed_audits} audit(s) closed in the year, the internal audit activity "
        f"{rating.lower() if rating != 'Not assessed' else 'has not yet been assessed'} with the Standards"
        + (f" (average conformance {q.average_score}%)." if q.average_score is not None else "."))
    doc.add_heading("2. Scope and approach", level=1)
    doc.add_paragraph("Each completed engagement was reviewed against a 12-point checklist mapped to the Standards "
                      "by a reviewer independent of the engagement team. Results are combined with performance "
                      "indicators and auditee feedback collected in the GPHA Internal Audit Workspace.")
    doc.add_heading("3. Conformance by area", level=1)
    weights = {"Conforms": 1.0, "Partially conforms": 0.5, "Does not conform": 0.0}
    per_item: dict[str, list[float]] = {}
    ids = [r["audit_id"] for r in q.reviews]
    for r in db.execute(select(QualityReview).where(QualityReview.audit_id.in_(ids or [0]))).scalars():
        for k, ans in json.loads(r.answers).items():
            if ans.get("answer") in weights:
                per_item.setdefault(k, []).append(weights[ans["answer"]])
    table(["Standard", "Checklist area", "Conformance"],
          [[s, qtext, f"{round(100 * sum(per_item[k]) / len(per_item[k]))}%" if per_item.get(k) else "-"]
           for k, s, qtext in CHECKLIST])
    doc.add_heading("4. Performance against targets", level=1)
    table(["Indicator", "Actual", "Target", "Status"],
          [[k.label, "-" if k.actual is None else f"{k.actual}{'%' if k.unit == '%' else ' ' + k.unit}",
            f"{'<=' if k.better == 'lower' else '>='} {k.target}{'%' if k.unit == '%' else ' ' + k.unit}",
            {"ok": "On target", "warn": "Near target", "bad": "Off target", "none": "No data"}[k.status]]
           for k in a.kpis])
    doc.add_heading("5. Stakeholder feedback", level=1)
    if q.survey_responses:
        doc.add_paragraph(f"{q.survey_responses} response(s) from {q.survey_invited} invited auditee(s); "
                          f"average satisfaction {q.survey_average} out of 5.")
        table(["Question", "Average (1-5)"], [[s["question"], s["average"]] for s in q.survey_by_question])
        for c in q.recent_comments[:5]:
            doc.add_paragraph(f'"{c}"', style="List Bullet")
    else:
        doc.add_paragraph("No survey responses were received in the year.")
    doc.add_heading("6. Areas for improvement", level=1)
    weak = [w for w in q.weakest_items if w["conformance_pct"] < 85]
    if weak:
        table(["Standard", "Area", "Conformance", "Improvement action (owner, date)"],
              [[w["standard"], w["question"], f"{w['conformance_pct']}%", ""] for w in weak])
    else:
        doc.add_paragraph("No checklist area fell below 85% conformance.")
    doc.add_heading("7. Approval", level=1)
    doc.add_paragraph("Chief Audit Executive: ____________________    Date: ____________\n\n"
                      "Presented to the Audit Committee on: ____________")
    buf = io.BytesIO()
    doc.save(buf)
    log_action(db, user_id=user.id, action=f"QAIP_REPORT fy={fy}", request=request)
    db.commit()
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="internal-assessment-FY{fy}.docx"'})
