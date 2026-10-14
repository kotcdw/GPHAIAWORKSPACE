"""Audit Committee reporting pack (IIA GIAS Domain III - communicating with the board)."""

from __future__ import annotations

import io
from collections import Counter
from datetime import datetime

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession

from app.core.appconfig import cfg
from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import staff_or_perm
from app.core.enums import CapaStatus, EngagementStatus, FindingStatus, ReportStatus, utcnow
from app.modules.audits.models import Audit, Finding
from app.modules.capa.models import Capa
from app.modules.followup.service import LEVEL_NAMES, overdue_actions
from app.modules.planning.models import AssurancePlan, PlanEngagement
from app.modules.reports.models import AuditReport
from app.modules.users.models import User

router = APIRouter(prefix="/api/committee", tags=["audit committee"])

Staff = Depends(staff_or_perm("governance.view"))
AGE_BUCKETS = ("0-30 days", "31-90 days", "91-180 days", "Over 180 days")


class PlanProgress(BaseModel):
    plan_ref: str | None = None
    title: str | None = None
    fiscal_year: int | None = None
    status: str | None = None
    engagements: int = 0
    by_stage: dict[str, int] = {}
    completed_pct: int = 0
    hours_planned: int = 0
    budget_hours: int = 0
    approved_at: datetime | None = None
    committee_meeting_date: datetime | None = None
    committee_reference: str | None = None


class IssuedReport(BaseModel):
    report_ref: str
    title: str
    rating: str | None
    issued_at: datetime | None
    audit_id: int


class OverdueRow(BaseModel):
    capa_ref: str
    title: str
    owner: str | None
    days_overdue: int
    escalated_to: str
    id: int


class CommitteeSummary(BaseModel):
    generated_at: datetime
    plan: PlanProgress
    audits_by_status: dict[str, int]
    reports_issued: list[IssuedReport]
    ratings: dict[str, int]
    open_findings_by_rating: dict[str, int]
    open_findings_by_age: dict[str, int]
    actions: dict[str, int]
    top_overdue: list[OverdueRow]
    plan_changes: list[dict] = []
    fiscal_years: list[int]


def _age_bucket(days: int) -> str:
    if days <= 30:
        return AGE_BUCKETS[0]
    if days <= 90:
        return AGE_BUCKETS[1]
    if days <= 180:
        return AGE_BUCKETS[2]
    return AGE_BUCKETS[3]


def build_summary(db: DBSession, fiscal_year: int | None = None) -> CommitteeSummary:
    now = utcnow()
    years = sorted(db.execute(select(AssurancePlan.fiscal_year)).scalars().all(), reverse=True)
    year = fiscal_year or (years[0] if years else now.year)
    plan = db.execute(select(AssurancePlan).where(AssurancePlan.fiscal_year == year)).scalar_one_or_none()

    progress = PlanProgress(fiscal_year=year)
    if plan is not None:
        engs = db.execute(select(PlanEngagement).where(PlanEngagement.plan_id == plan.id)).scalars().all()
        stages: Counter = Counter()
        for e in engs:
            if e.status == EngagementStatus.Cancelled:
                stages["Cancelled"] += 1
            elif e.audit_id is None:
                stages["Not started"] += 1
            else:
                audit = db.get(Audit, e.audit_id)
                stages["Completed" if audit and audit.status.value == "Closed" else
                       (audit.status.value if audit else "Not started")] += 1
        live = len(engs) - stages["Cancelled"]
        progress = PlanProgress(
            plan_ref=plan.plan_ref, title=plan.title, fiscal_year=plan.fiscal_year, status=plan.status.value,
            engagements=len(engs), by_stage=dict(stages),
            completed_pct=round(100 * stages["Completed"] / live) if live else 0,
            hours_planned=sum(e.budget_hours for e in engs if e.status != EngagementStatus.Cancelled),
            budget_hours=plan.total_budget_hours, approved_at=plan.approved_at,
            committee_meeting_date=plan.committee_meeting_date, committee_reference=plan.committee_reference,
        )

    audits = db.execute(select(Audit)).scalars().all()
    reports = db.execute(
        select(AuditReport).where(AuditReport.status == ReportStatus.Issued).order_by(AuditReport.issued_at.desc())
    ).scalars().all()
    reports = [r for r in reports if r.issued_at and r.issued_at.year == year] or reports[:10]

    open_findings = db.execute(select(Finding).where(Finding.status != FindingStatus.Closed)).scalars().all()
    by_age: Counter = Counter({b: 0 for b in AGE_BUCKETS})
    for f in open_findings:
        by_age[_age_bucket((now - f.created_at).days)] += 1

    capas = db.execute(select(Capa)).scalars().all()
    overdue = overdue_actions(db, now)
    actions = {
        "open": sum(1 for c in capas if c.status != CapaStatus.Closed),
        "overdue": len(overdue),
        "awaiting_validation": sum(1 for c in capas if c.status == CapaStatus.Review and c.implemented_at
                                   and not c.validation_result),
        "closed": sum(1 for c in capas if c.status == CapaStatus.Closed),
        "reopened": sum(1 for c in capas if (c.reopen_count or 0) > 0),
        "with_cae": sum(1 for o in overdue if o.due_level >= 3),
        "for_committee": sum(1 for o in overdue if o.due_level >= 4),
    }
    changes: list[dict] = []
    if plan is not None:
        from app.modules.planning.models import PlanChange

        changes = [{"date": c.changed_at.date().isoformat(), "type": c.change_type, "description": c.description,
                    "reason": c.reason}
                   for c in db.execute(select(PlanChange).where(PlanChange.plan_id == plan.id,
                                                                PlanChange.change_type != "Committee")
                                       .order_by(PlanChange.changed_at)).scalars()]
    return CommitteeSummary(
        generated_at=now,
        plan=progress,
        audits_by_status=dict(Counter(a.status.value for a in audits)),
        reports_issued=[IssuedReport(report_ref=r.report_ref, title=r.title, audit_id=r.audit_id,
                                     rating=r.overall_rating.value if r.overall_rating else None,
                                     issued_at=r.issued_at) for r in reports],
        ratings=dict(Counter(r.overall_rating.value for r in reports if r.overall_rating)),
        open_findings_by_rating={k: sum(1 for f in open_findings if f.risk_rating.value == k)
                                 for k in ("High", "Medium", "Low")},
        open_findings_by_age=dict(by_age),
        actions=actions,
        top_overdue=[OverdueRow(capa_ref=o.capa.capa_ref, title=o.capa.title, owner=o.capa.owner_email,
                                days_overdue=o.days_overdue, escalated_to=LEVEL_NAMES[o.due_level], id=o.capa.id)
                     for o in sorted(overdue, key=lambda o: -o.days_overdue)[:10]],
        plan_changes=changes,
        fiscal_years=years,
    )


@router.get("/summary", response_model=CommitteeSummary)
def committee_summary(fiscal_year: int | None = None, db: DBSession = Depends(get_db),
                      _: User = Staff) -> CommitteeSummary:
    return build_summary(db, fiscal_year)


@router.get("/pack.docx")
def committee_pack(request: Request, fiscal_year: int | None = None, db: DBSession = Depends(get_db),
                   user: User = Staff) -> Response:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    s = build_summary(db, fiscal_year)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.add_run(f"{cfg('org_name').upper()}\n{cfg('department_name').upper()}").bold = True
    doc.add_heading(f"Report to the Audit Committee - FY{s.plan.fiscal_year}", level=0)
    doc.add_paragraph(f"Prepared {s.generated_at:%d %B %Y} from the GPHA Internal Audit Workspace.")

    def table(headers: list[str], rows: list[list[str]]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for cell, text in zip(t.rows[0].cells, headers, strict=True):
            cell.text = text
        for row in rows:
            for cell, text in zip(t.add_row().cells, row, strict=True):
                cell.text = str(text)

    doc.add_heading("1. Delivery of the Internal Audit Plan", level=1)
    p = s.plan
    if p.plan_ref:
        doc.add_paragraph(f"{p.plan_ref} {p.title} ({p.status}): {p.engagements} engagements, "
                          f"{p.completed_pct}% completed; {p.hours_planned} of {p.budget_hours} budget hours planned.")
        if p.committee_meeting_date:
            doc.add_paragraph(f"Approved by the Audit Committee on {p.committee_meeting_date:%d %B %Y} "
                              f"(ref {p.committee_reference}).")
        else:
            doc.add_paragraph("The plan has not yet been recorded as approved by the Audit Committee.")
        table(["Stage", "Engagements"], [[k, v] for k, v in p.by_stage.items()])
    else:
        doc.add_paragraph("No Assurance Plan recorded for this year.")
    doc.add_heading("2. Reports issued", level=1)
    if s.reports_issued:
        table(["Report", "Title", "Overall rating", "Issued"],
              [[r.report_ref, r.title, r.rating or "-", f"{r.issued_at:%d %b %Y}" if r.issued_at else "-"]
               for r in s.reports_issued])
    else:
        doc.add_paragraph("No reports issued in the period.")
    doc.add_heading("3. Open findings", level=1)
    table(["Risk rating", "Open findings"], [[k, v] for k, v in s.open_findings_by_rating.items()])
    doc.add_paragraph("")
    table(["Age", "Open findings"], [[k, v] for k, v in s.open_findings_by_age.items()])
    doc.add_heading("4. Implementation of agreed actions", level=1)
    a = s.actions
    doc.add_paragraph(f"{a['open']} open, {a['overdue']} overdue ({a['with_cae']} escalated to the CAE, "
                      f"{a['for_committee']} reported to this Committee), "
                      f"{a['awaiting_validation']} awaiting auditor validation, {a['closed']} closed, "
                      f"{a['reopened']} reopened after validation.")
    if s.top_overdue:
        table(["Action", "Description", "Owner", "Days overdue", "Escalated to"],
              [[o.capa_ref, o.title, o.owner or "-", o.days_overdue, o.escalated_to] for o in s.top_overdue])
    doc.add_heading("5. Changes to the approved plan", level=1)
    if s.plan_changes:
        table(["Date", "Change", "Description", "Reason"],
              [[c["date"], c["type"], c["description"], c["reason"]] for c in s.plan_changes])
    else:
        doc.add_paragraph("No changes since the plan was approved.")
    buf = io.BytesIO()
    doc.save(buf)
    log_action(db, user_id=user.id, action=f"COMMITTEE_PACK fy={p.fiscal_year}", request=request)
    db.commit()
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="Audit-Committee-Pack-FY{p.fiscal_year}.docx"'},
    )
