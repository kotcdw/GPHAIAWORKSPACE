"""Audit report content, snapshots and exports (Word and printable HTML)."""

from __future__ import annotations

import html
import io
import json
from datetime import datetime

from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession

from app.core.appconfig import cfg
from app.core.enums import RiskRating
from app.modules.audits.models import Audit, Finding
from app.modules.capa.models import Capa
from app.modules.reports.models import AuditReport
from app.modules.users.models import User

_RISK_ORDER = {RiskRating.High.value: 0, RiskRating.Medium.value: 1, RiskRating.Low.value: 2}


def _name(db: DBSession, user_id: int | None) -> str | None:
    if user_id is None:
        return None
    user = db.get(User, user_id)
    return (user.full_name or user.username) if user else None


def _fmt(dt: datetime | None) -> str:
    return dt.strftime("%d %b %Y") if dt else ""


def live_findings(db: DBSession, audit_id: int) -> list[dict]:
    findings = db.execute(select(Finding).where(Finding.audit_id == audit_id)).scalars().all()
    out = []
    for f in findings:
        actions = db.execute(
            select(Capa).where(Capa.source_type == "finding", Capa.source_id == f.id)
        ).scalars().all()
        out.append(
            {
                "id": f.id,
                "title": f.title,
                "risk_rating": f.risk_rating.value,
                "status": f.status.value,
                "criteria": f.criteria or "",
                "condition": f.condition or f.description or "",
                "cause": f.cause or "",
                "effect": f.effect or "",
                "recommendation": f.recommendation or "",
                "management_response": f.management_response or "",
                "owner": f.assigned_owner_email or "",
                "actions": [
                    {"ref": c.capa_ref, "owner": c.owner_email or "", "due": _fmt(c.due_date), "status": c.status.value}
                    for c in actions
                ],
            }
        )
    out.sort(key=lambda d: (_RISK_ORDER.get(d["risk_rating"], 9), d["id"]))
    return out


def signoffs(db: DBSession, report: AuditReport) -> list[dict]:
    return [
        {"role": "Prepared by", "name": _name(db, report.created_by), "at": report.created_at},
        {"role": "Reviewed by (Audit Manager)", "name": _name(db, report.manager_reviewed_by),
         "at": report.manager_reviewed_at},
        {"role": "Approved by (Chief Audit Executive)", "name": _name(db, report.cae_approved_by),
         "at": report.cae_approved_at},
    ]


def take_snapshot(db: DBSession, report: AuditReport) -> str:
    return json.dumps({"findings": live_findings(db, report.audit_id)}, default=str)


def content(db: DBSession, report: AuditReport) -> dict:
    """Everything an export needs; issued reports read their frozen snapshot."""
    audit = db.get(Audit, report.audit_id)
    if report.issued_snapshot:
        findings = json.loads(report.issued_snapshot).get("findings", [])
    else:
        findings = live_findings(db, report.audit_id)
    counts = {k: sum(1 for f in findings if f["risk_rating"] == k) for k in ("High", "Medium", "Low")}
    return {
        "report": report,
        "audit": audit,
        "audit_ref": f"AUD-{audit.id:04d}" if audit else "",
        "findings": findings,
        "counts": counts,
        "signoffs": signoffs(db, report),
        "draft": report.status.value != "Issued",
    }


# ------------------------------------------------------------------ Word export


def to_docx(db: DBSession, report: AuditReport) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    data = content(db, report)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(f"{cfg('org_name').upper()}\n{cfg('department_name').upper()}")
    run.bold = True
    run.font.size = Pt(13)
    if data["draft"]:
        mark = doc.add_paragraph()
        mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = mark.add_run(f"DRAFT - {report.status.value.upper()} - NOT FOR DISTRIBUTION")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    doc.add_heading(report.title, level=0)
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light Grid Accent 1"
    for label, value in (
        ("Report reference", report.report_ref),
        ("Audit reference", data["audit_ref"]),
        ("Overall rating", report.overall_rating.value if report.overall_rating else "Not rated"),
        ("Status", report.status.value),
        ("Date issued", _fmt(report.issued_at) or "Not issued"),
        ("Findings", f"{len(data['findings'])} (High {data['counts']['High']}, Medium {data['counts']['Medium']}, "
                     f"Low {data['counts']['Low']})"),
    ):
        cells = meta.add_row().cells
        cells[0].text = label
        cells[1].text = value

    for heading, text in (
        ("1. Executive summary", report.executive_summary),
        ("2. Background", report.background),
        ("3. Scope and objectives", report.scope_objectives),
        ("4. Methodology", report.methodology),
        ("5. Overall opinion", report.overall_opinion),
    ):
        doc.add_heading(heading, level=1)
        for para in (text or "Not provided.").split("\n"):
            doc.add_paragraph(para)

    doc.add_heading("6. Detailed findings", level=1)
    if not data["findings"]:
        doc.add_paragraph("No findings were raised in this engagement.")
    for i, f in enumerate(data["findings"], start=1):
        doc.add_heading(f"6.{i} {f['title']} ({f['risk_rating']} risk)", level=2)
        for label, key in (
            ("Criteria", "criteria"),
            ("Condition", "condition"),
            ("Cause", "cause"),
            ("Effect", "effect"),
            ("Recommendation", "recommendation"),
            ("Management response", "management_response"),
        ):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(f[key] or "Not provided.")
        if f["actions"]:
            p = doc.add_paragraph()
            p.add_run("Agreed actions: ").bold = True
            p.add_run("; ".join(f"{a['ref']} (owner {a['owner'] or 'n/a'}, due {a['due'] or 'n/a'}, {a['status']})"
                                for a in f["actions"]))

    doc.add_heading("7. Management comments", level=1)
    doc.add_paragraph(report.management_comments or "None received.")
    doc.add_heading("8. Distribution", level=1)
    doc.add_paragraph(report.distribution_list or "Not specified.")

    doc.add_heading("Sign-off", level=1)
    sign = doc.add_table(rows=1, cols=3)
    sign.style = "Light Grid Accent 1"
    hdr = sign.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Role", "Name", "Date"
    for s in data["signoffs"]:
        cells = sign.add_row().cells
        cells[0].text = s["role"]
        cells[1].text = s["name"] or "Pending"
        cells[2].text = _fmt(s["at"]) if s["name"] else ""

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ printable HTML (Save as PDF)


def to_html(db: DBSession, report: AuditReport) -> str:
    data = content(db, report)
    e = html.escape

    def paras(text: str) -> str:
        return "".join(f"<p>{e(line)}</p>" for line in (text or "Not provided.").split("\n") if line.strip()) or \
            "<p>Not provided.</p>"

    rows = "".join(
        f"<tr><td>{e(f['title'])}</td><td class='r-{f['risk_rating'].lower()}'>{e(f['risk_rating'])}</td>"
        f"<td>{e(f['status'])}</td><td>{e(', '.join(a['ref'] for a in f['actions']) or '-')}</td></tr>"
        for f in data["findings"]
    )
    details = ""
    for i, f in enumerate(data["findings"], start=1):
        actions = "; ".join(
            f"{a['ref']} (owner {a['owner'] or 'n/a'}, due {a['due'] or 'n/a'}, {a['status']})" for a in f["actions"]
        )
        details += (
            f"<section class='finding'><h3>6.{i} {e(f['title'])} <span class='r-{f['risk_rating'].lower()}'>"
            f"{e(f['risk_rating'])} risk</span></h3><dl>"
            + "".join(
                f"<dt>{label}</dt><dd>{e(f[key]) or 'Not provided.'}</dd>"
                for label, key in (
                    ("Criteria", "criteria"), ("Condition", "condition"), ("Cause", "cause"), ("Effect", "effect"),
                    ("Recommendation", "recommendation"), ("Management response", "management_response"),
                )
            )
            + (f"<dt>Agreed actions</dt><dd>{e(actions)}</dd>" if actions else "")
            + "</dl></section>"
        )
    signs = "".join(
        f"<tr><td>{e(s['role'])}</td><td>{e(s['name'] or 'Pending')}</td><td>{_fmt(s['at']) if s['name'] else ''}</td></tr>"
        for s in data["signoffs"]
    )
    rating = report.overall_rating.value if report.overall_rating else "Not rated"
    banner = (f"<div class='draft'>DRAFT - {e(report.status.value)} - not for distribution</div>"
              if data["draft"] else "")
    return f"""<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(report.report_ref)} {e(report.title)}</title>
<style>
body{{font-family:Calibri,'Segoe UI',Arial,sans-serif;color:#1d2433;max-width:820px;margin:32px auto;padding:0 24px;line-height:1.5}}
header{{text-align:center;border-bottom:2px solid #1f4e9c;padding-bottom:12px;margin-bottom:20px}}
header b{{display:block;letter-spacing:.04em}} h1{{font-size:24px;margin:18px 0 8px}} h2{{font-size:17px;color:#1f4e9c;margin-top:26px;border-bottom:1px solid #d6dde8}}
h3{{font-size:15px;margin:18px 0 6px}} table{{border-collapse:collapse;width:100%;margin:8px 0}} td,th{{border:1px solid #cfd7e3;padding:6px 8px;text-align:left;vertical-align:top}}
th{{background:#eef2f8}} dl{{margin:0}} dt{{font-weight:700;margin-top:6px}} dd{{margin:0 0 4px 0}}
.r-high{{color:#b42318;font-weight:700}} .r-medium{{color:#b54708;font-weight:700}} .r-low{{color:#067647;font-weight:700}}
.draft{{border:2px dashed #c0392b;color:#c0392b;text-align:center;font-weight:700;padding:6px;margin-bottom:12px}}
.finding{{page-break-inside:avoid}} .noprint{{text-align:right}} @media print{{.noprint{{display:none}} body{{margin:0}}}}
</style></head><body>
<div class='noprint'><button onclick='window.print()'>Print / Save as PDF</button></div>
<header><b>{e(cfg('org_name').upper())}</b><b>{e(cfg('department_name').upper())}</b></header>
{banner}<h1>{e(report.title)}</h1>
<table><tr><th>Report reference</th><td>{e(report.report_ref)}</td><th>Audit reference</th><td>{e(data['audit_ref'])}</td></tr>
<tr><th>Overall rating</th><td>{e(rating)}</td><th>Date issued</th><td>{_fmt(report.issued_at) or 'Not issued'}</td></tr>
<tr><th>Findings</th><td colspan='3'>{len(data['findings'])} - High {data['counts']['High']}, Medium {data['counts']['Medium']}, Low {data['counts']['Low']}</td></tr></table>
<h2>1. Executive summary</h2>{paras(report.executive_summary)}
<h2>2. Background</h2>{paras(report.background)}
<h2>3. Scope and objectives</h2>{paras(report.scope_objectives)}
<h2>4. Methodology</h2>{paras(report.methodology)}
<h2>5. Overall opinion</h2>{paras(report.overall_opinion)}
<h2>6. Findings</h2>{"<table><tr><th>Finding</th><th>Risk</th><th>Status</th><th>Actions</th></tr>" + rows + "</table>" if rows else "<p>No findings were raised in this engagement.</p>"}
{details}
<h2>7. Management comments</h2>{paras(report.management_comments or 'None received.')}
<h2>8. Distribution</h2>{paras(report.distribution_list or 'Not specified.')}
<h2>Sign-off</h2><table><tr><th>Role</th><th>Name</th><th>Date</th></tr>{signs}</table>
</body></html>"""
