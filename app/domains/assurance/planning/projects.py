"""Assurance plan projects (Diligent: Assurance Plans > Add to plan, plan overview, coverage, risks, resources).

* What a planned engagement brings into its audit: audit type, an audit programme, or (roll forward) the risk &
  control structure of an earlier audit. When none is chosen, the auditable entity's linked risks and controls are
  scoped in. ``prepare_started_audit`` runs when the engagement is started.
* Roll forward: plan an engagement from an earlier audit.
* Programme overview of one plan: key metrics, projects with progress, strategic risks addressed and assignees,
  coverage by entity category, strategic risks (Risk Manager risks linked to the plan's entities or audit risks,
  with the strategic objective they threaten), and resources.
* Registers across the plan's audits: objectives and narratives (they are kept inside each audit), and importing a
  narrative from a document into an audit objective.

A plan holds at most ``plan_max_projects`` live engagements (default 100).
"""

from __future__ import annotations

import io
from collections import defaultdict

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session as DBSession

from app.core.appconfig import cfg
from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import require_role
from app.core.enums import AuditStatus, EngagementStatus, PlanStatus, Role
from app.core.errors import field_error
from app.core.refs import next_reference
from app.domains.assurance.planning.models import AssurancePlan, EntityUnit, PlanEngagement
from app.domains.audit.audits.models import Audit, AuditTeamMember
from app.platform.users.models import User

router = APIRouter(tags=["planning"])
Staff = Depends(require_role(Role.Admin, Role.Auditor))


def _plan(db: DBSession, plan_id: int) -> AssurancePlan:
    plan = db.get(AssurancePlan, plan_id)
    if plan is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Assurance plan not found")
    return plan


def _conflict(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_409_CONFLICT, detail=detail)


def live(plan: AssurancePlan) -> list[PlanEngagement]:
    return [e for e in plan.engagements if e.status != EngagementStatus.Cancelled]


def max_projects() -> int:
    try:
        return int(cfg("plan_max_projects"))
    except (KeyError, TypeError, ValueError):
        return 100


def ensure_room(plan: AssurancePlan, adding: int = 1) -> None:
    limit = max_projects()
    if len(live(plan)) + adding > limit:
        raise _conflict(f"A plan holds at most {limit} projects ({len(live(plan))} already planned)")


# ------------------------------------------------------------------ engagement set-up


def check_setup(db: DBSession, audit_type: str | None, program_id: int | None, source_audit_id: int | None) -> None:
    """Validate what the engagement brings into its audit (0 / "" clear a value on update)."""
    from app.domains.audit.programs.models import AuditProgram

    if audit_type:
        if audit_type not in (cfg("audit_types") or []):
            raise field_error("audit_type", f"must be one of: {', '.join(cfg('audit_types') or [])}")
    if program_id:
        p = db.get(AuditProgram, program_id)
        if p is None or p.status != "Active":
            raise field_error("program_id", "choose an active audit programme from the Project Library")
    if source_audit_id:
        if db.get(Audit, source_audit_id) is None:
            raise field_error("source_audit_id", "the audit to roll forward does not exist")
    if program_id and source_audit_id:
        raise field_error("program_id", "choose an audit programme or roll forward an earlier audit, not both")


def apply_setup(eng: PlanEngagement, sent: dict) -> None:
    for key in ("audit_type", "program_id", "source_audit_id"):
        if key in sent and sent[key] is not None:
            setattr(eng, key, sent[key] or None)


def setup_out(db: DBSession, eng: PlanEngagement, out) -> None:  # noqa: ANN001
    from app.domains.audit.programs.models import AuditProgram

    if eng.program_id:
        p = db.get(AuditProgram, eng.program_id)
        out.program_name = f"{p.prog_ref} {p.name}" if p else None
    if eng.source_audit_id:
        out.source_audit_ref = f"AUD-{eng.source_audit_id:04d}"


def _roll_forward(db: DBSession, source: Audit, audit: Audit, user: User) -> dict:
    """Copy the earlier audit's objectives, risks, controls with their test plan, narratives and procedures (not
    its results, tests or sign-offs) into the new audit."""
    from app.domains.audit.rcm.models import (
        Control,
        Narrative,
        NarrativeControl,
        Objective,
        Procedure,
        ProcedureRisk,
        Risk,
        RiskControl,
    )

    counts = {"objectives": 0, "risks": 0, "controls": 0, "narratives": 0, "procedures": 0}
    for o in db.execute(select(Objective).where(Objective.audit_id == source.id).order_by(Objective.id)).scalars():
        no = Objective(audit_id=audit.id, ref=next_reference(db, "OBJ"), title=o.title, description=o.description,
                       created_by=user.id)
        db.add(no)
        db.flush()
        counts["objectives"] += 1
        risks: dict[int, int] = {}
        for r in o.risks:
            nr = Risk(objective_id=no.id, ref=next_reference(db, "RSK"), title=r.title, description=r.description,
                      likelihood=r.likelihood, impact=r.impact, enterprise_risk_id=r.enterprise_risk_id,
                      created_by=user.id)
            db.add(nr)
            db.flush()
            risks[r.id] = nr.id
            counts["risks"] += 1
        controls: dict[int, int] = {}
        for c in o.controls:
            nc = Control(objective_id=no.id, ref=next_reference(db, "CTL"), title=c.title, description=c.description,
                         control_type=c.control_type, nature=c.nature, frequency=c.frequency,
                         owner_email=c.owner_email, is_key=c.is_key, library_control_id=c.library_control_id,
                         test_method=c.test_method, total_sample_size=c.total_sample_size, test_steps=c.test_steps,
                         created_by=user.id)
            db.add(nc)
            db.flush()
            controls[c.id] = nc.id
            counts["controls"] += 1
            for r in c.risks:
                if r.id in risks:
                    db.add(RiskControl(risk_id=risks[r.id], control_id=nc.id))
        for n in db.execute(select(Narrative).where(Narrative.objective_id == o.id)
                            .order_by(Narrative.position)).scalars():
            nn = Narrative(objective_id=no.id, ref=next_reference(db, "NAR"), title=n.title,
                           description=n.description, position=n.position, custom_fields="{}", created_by=user.id)
            db.add(nn)
            db.flush()
            counts["narratives"] += 1
            for c in n.controls:
                if c.id in controls:
                    db.add(NarrativeControl(narrative_id=nn.id, control_id=controls[c.id]))
        for p in db.execute(select(Procedure).where(Procedure.objective_id == o.id).order_by(Procedure.id)).scalars():
            np_ = Procedure(objective_id=no.id, ref=next_reference(db, "PRC"), title=p.title, steps=p.steps,
                            status="Not Started", results="", review_comment="", created_by=user.id)
            db.add(np_)
            db.flush()
            counts["procedures"] += 1
            for r in p.risks:
                if r.id in risks:
                    db.add(ProcedureRisk(procedure_id=np_.id, risk_id=risks[r.id]))
    db.flush()
    return counts


def prepare_started_audit(db: DBSession, eng: PlanEngagement, audit: Audit, user: User) -> str:
    """Bring the engagement's set-up into the audit just created from it. Returns a note for the history."""
    from app.domains.audit.fieldwork.router import _scope_plan, default_workflow, scope_into
    from app.domains.audit.programs.models import AuditProgram
    from app.domains.audit.programs.router import copy_program

    if eng.audit_type:
        audit.audit_type = eng.audit_type
    source = db.get(Audit, eng.source_audit_id) if eng.source_audit_id else None
    audit.workflow = source.workflow if source else default_workflow(audit.audit_type)
    db.flush()
    if source is not None:
        c = _roll_forward(db, source, audit, user)
        return f"rolled forward from AUD-{source.id:04d} ({c['objectives']} objective(s), {c['controls']} control(s))"
    program = db.get(AuditProgram, eng.program_id) if eng.program_id else None
    if program is not None and program.status == "Active" and program.step_list:
        copy_program(db, audit, program, user)
        return f"audit programme {program.prog_ref} applied"
    if eng.entity_unit_id and _scope_plan(db, audit):
        c = scope_into(db, audit, user)
        return f"scoped from the auditable entity ({c['risks']} risk(s), {c['controls'] or c['procedures']} item(s))"
    return ""


# ------------------------------------------------------------------ roll forward


@router.get("/api/assurance-plans/{plan_id}/rollforward-candidates")
def rollforward_candidates(plan_id: int, q: str | None = Query(None, max_length=100),
                           db: DBSession = Depends(get_db), user: User = Staff) -> list[dict]:
    """Earlier audits that can be rolled forward into this plan (any audit you can see that is not on it)."""
    from app.core.access import visible

    plan = _plan(db, plan_id)
    audits = visible(db, user, "audit", list(db.execute(select(Audit).order_by(Audit.id.desc())).scalars()))
    plans = {p.id: p for p in db.execute(select(AssurancePlan)).scalars()}
    out = []
    for a in audits:
        if a.plan_id == plan.id:
            continue
        ref = f"AUD-{a.id:04d}"
        if q and q.lower() not in f"{ref} {a.title} {a.audit_type or ''}".lower():
            continue
        p = plans.get(a.plan_id)
        out.append({"id": a.id, "ref": ref, "title": a.title, "audit_type": a.audit_type, "status": a.status.value,
                    "plan": f"FY{p.fiscal_year} · {p.plan_ref}" if p else None, "entity_unit_id": a.entity_unit_id})
        if len(out) >= 200:
            break
    return out


class RollforwardIn(BaseModel):
    audit_id: int
    title: str = Field(min_length=1, max_length=255)
    planned_quarter: str = Field("Q1", pattern="^Q[1-4]$")
    budget_hours: int = Field(0, ge=0, le=10_000)
    lead_auditor_id: int | None = None
    site_entity_id: int | None = None


@router.post("/api/assurance-plans/{plan_id}/rollforward", status_code=status.HTTP_201_CREATED)
def rollforward(plan_id: int, payload: RollforwardIn, request: Request, db: DBSession = Depends(get_db),
                user: User = Staff) -> dict:
    """Plan an engagement from an earlier audit: same entity, audit type and site; when the audit starts it
    receives the earlier audit's objectives, risks, controls, test plan and narratives."""
    from app.core.deps import ensure_perm
    from app.domains.assurance.planning.router import _check_refs, _engagement_out, _history, check_site

    ensure_perm(user, "plan.prepare", "plan.amend")
    plan = _plan(db, plan_id)
    if plan.status != PlanStatus.Draft:
        raise _conflict("Engagements are added to a draft plan; use a change request for an approved plan")
    ensure_room(plan)
    source = db.get(Audit, payload.audit_id)
    if source is None:
        raise field_error("audit_id", "choose the audit to roll forward")
    _check_refs(db, None, payload.lead_auditor_id)
    check_site(db, payload.site_entity_id)
    entity = db.get(EntityUnit, source.entity_unit_id) if source.entity_unit_id else None
    eng = PlanEngagement(plan_id=plan.id, title=payload.title.strip(), planned_quarter=payload.planned_quarter,
                         budget_hours=payload.budget_hours, lead_auditor_id=payload.lead_auditor_id,
                         site_entity_id=payload.site_entity_id or source.port_entity_id,
                         entity_unit_id=entity.id if entity is not None and entity.is_active else None,
                         audit_type=source.audit_type, source_audit_id=source.id,
                         rationale=f"Rolled forward from AUD-{source.id:04d} {source.title}.",
                         status=EngagementStatus.Planned, created_by=user.id)
    db.add(eng)
    db.flush()
    log_action(db, user_id=user.id, action=f"ENGAGEMENT_ROLLFORWARD id={eng.id} from=AUD-{source.id:04d}",
               request=request)
    _history(db, plan.id, user, f"ENGAGEMENT ROLLED FORWARD #{eng.id} {eng.title} from AUD-{source.id:04d}", request)
    db.commit()
    db.refresh(eng)
    return _engagement_out(eng).model_dump(mode="json")


# ------------------------------------------------------------------ programme overview


def _strategic_risks(db: DBSession, plan: AssurancePlan, audits: dict[int, Audit]) -> dict[int, set[int]]:
    """Enterprise (strategic) risk id -> engagement ids that address it: linked to the engagement's entity (Process
    library / Risk Manager links) or to a risk inside its audit."""
    from app.domains.audit.rcm.models import Objective, Risk
    from app.domains.grc.links import linked_ids

    out: dict[int, set[int]] = defaultdict(set)
    for e in live(plan):
        if e.entity_unit_id:
            for rid in linked_ids(db, "unit", e.entity_unit_id, "risk"):
                out[rid].add(e.id)
        if e.audit_id and e.audit_id in audits:
            for (rid,) in db.execute(select(Risk.enterprise_risk_id).join(Objective, Objective.id == Risk.objective_id)
                                     .where(Objective.audit_id == e.audit_id,
                                            Risk.enterprise_risk_id.is_not(None))).all():
                out[rid].add(e.id)
    return out


@router.get("/api/assurance-plans/{plan_id}/program")
def program_overview(plan_id: int, db: DBSession = Depends(get_db), user: User = Staff) -> dict:
    from app.domains.assurance.planning.coverage import plan_coverage
    from app.domains.grc.risk.models import EnterpriseRisk

    plan = _plan(db, plan_id)
    cov = plan_coverage(plan_id, None, False, db, user)
    audits = {a.id: a for a in db.execute(select(Audit).where(Audit.id.in_(
        [e.audit_id for e in plan.engagements if e.audit_id] or [0]))).scalars()}
    team = dict(db.execute(select(AuditTeamMember.audit_id, func.count()).where(
        AuditTeamMember.audit_id.in_(list(audits) or [0])).group_by(AuditTeamMember.audit_id)).all())
    risk_map = _strategic_risks(db, plan, audits)
    by_eng: dict[int, int] = defaultdict(int)
    for engs in risk_map.values():
        for eid in engs:
            by_eng[eid] += 1
    progress = {r["engagement_id"]: r["progress"] for r in cov["engagements"]}
    entity_names = {u.id: u.name for u in db.execute(select(EntityUnit).where(EntityUnit.id.in_(
        [e.entity_unit_id for e in plan.engagements if e.entity_unit_id] or [0]))).scalars()}
    projects = []
    for r in cov["engagements"]:
        e = next(x for x in plan.engagements if x.id == r["engagement_id"])
        a = audits.get(e.audit_id) if e.audit_id else None
        projects.append({
            **r, "audit_type": (a.audit_type if a else e.audit_type), "entity": entity_names.get(e.entity_unit_id),
            "phase": ("Not started" if a is None else {"Planned": "Planning", "In Progress": "Fieldwork",
                                                       "Review": "Reporting", "Closed": "Completed"}[a.status.value]),
            "strategic_risks": by_eng.get(e.id, 0),
            "assignees": team.get(e.audit_id, 0) if a else (1 if e.lead_auditor_id else 0),
            "rolled_from": f"AUD-{e.source_audit_id:04d}" if e.source_audit_id else None,
        })
    # Coverage by entity category (location), over all active auditable entities.
    cats: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for u in cov["entities"]:
        c = cats[u["location"]]
        c[0] += 1
        c[1] += 1 if u["covered"] else 0
    total, covered = len(cov["entities"]), sum(1 for u in cov["entities"] if u["covered"])
    risks = []
    for rid, engs in risk_map.items():
        er = db.get(EnterpriseRisk, rid)
        if er is None or er.status == "Closed":
            continue
        score = er.residual_likelihood * er.residual_impact
        vals = [progress.get(eid, 0) for eid in engs]
        objective = None
        if er.objective_id:
            from app.domains.grc.risk.models import StrategicObjective

            so = db.get(StrategicObjective, er.objective_id)
            objective = so.title if so else None
        risks.append({"id": er.id, "ref": er.ref, "title": er.title, "category": er.category,
                      "objective": objective, "impact": er.residual_impact, "likelihood": er.residual_likelihood,
                      "score": score, "projects": len(engs), "progress": round(sum(vals) / len(vals)) if vals else 0})
    risks.sort(key=lambda x: (-x["score"], x["ref"]))
    return {
        "plan_id": plan.id,
        "metrics": {"projects": len(live(plan)), "max_projects": max_projects(), "strategic_risks": len(risks),
                    "coverage_pct": round(100 * covered / total) if total else 0, "entities": total,
                    "covered": covered, "hours": cov["hours"], "overall_progress": cov["overall_progress"]},
        "projects": projects,
        "categories": [{"name": k, "entities": v[0], "covered": v[1], "pct": round(100 * v[1] / v[0]) if v[0] else 0}
                       for k, v in sorted(cats.items())],
        "strategic_risks": risks,
        "not_covered_high": cov["not_covered_high"],
    }


# ------------------------------------------------------------------ objectives and narratives across the plan


def _plan_audits(db: DBSession, plan: AssurancePlan) -> dict[int, tuple[PlanEngagement, Audit]]:
    out = {}
    for e in live(plan):
        if e.audit_id:
            a = db.get(Audit, e.audit_id)
            if a is not None:
                out[a.id] = (e, a)
    return out


@router.get("/api/assurance-plans/{plan_id}/objectives")
def plan_objectives(plan_id: int, db: DBSession = Depends(get_db), _: User = Staff) -> list[dict]:
    """Every objective of the plan's audits, with its lead, target date and counts."""
    from app.domains.audit.rcm.models import Narrative, Objective

    plan = _plan(db, plan_id)
    pa = _plan_audits(db, plan)
    leads = {u.id: u.full_name or u.username for u in db.execute(select(User).where(User.id.in_(
        [e.lead_auditor_id for e, _ in pa.values() if e.lead_auditor_id] or [0]))).scalars()}
    out = []
    for o in db.execute(select(Objective).where(Objective.audit_id.in_(list(pa) or [0]))
                        .order_by(Objective.audit_id, Objective.id)).scalars():
        e, a = pa[o.audit_id]
        n = db.execute(select(func.count()).select_from(Narrative).where(Narrative.objective_id == o.id)).scalar_one()
        out.append({"id": o.id, "ref": o.ref, "title": o.title, "description": o.description,
                    "audit_id": a.id, "audit_ref": f"AUD-{a.id:04d}", "audit": a.title, "workflow": a.workflow,
                    "lead": leads.get(e.lead_auditor_id), "target_date": e.planned_end.isoformat()
                    if e.planned_end else None, "quarter": e.planned_quarter, "risks": len(o.risks),
                    "controls": len(o.controls), "narratives": n, "open": a.status != AuditStatus.Closed})
    return out


@router.get("/api/assurance-plans/{plan_id}/narratives")
def plan_narratives(plan_id: int, db: DBSession = Depends(get_db), _: User = Staff) -> list[dict]:
    from app.domains.audit.rcm.models import Narrative, Objective

    plan = _plan(db, plan_id)
    pa = _plan_audits(db, plan)
    rows = db.execute(select(Narrative, Objective).join(Objective, Objective.id == Narrative.objective_id)
                      .where(Objective.audit_id.in_(list(pa) or [0]))
                      .order_by(Objective.audit_id, Narrative.position)).all()
    return [{"id": n.id, "ref": n.ref, "title": n.title, "length": len(n.description or ""),
             "objective_id": o.id, "objective": f"{o.ref} {o.title}", "audit_id": o.audit_id,
             "audit": f"AUD-{o.audit_id:04d} {pa[o.audit_id][1].title}", "controls": len(n.controls),
             "updated_at": n.updated_at.isoformat()} for n, o in rows]


def _document_text(data: bytes, suffix: str) -> str:
    if suffix == ".txt":
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    raise field_error("file", "import a .docx, .xlsx or .txt narrative")


@router.post("/api/objectives/{objective_id}/narratives/import", status_code=status.HTTP_201_CREATED)
async def import_narrative(objective_id: int, request: Request, file: UploadFile = File(...),
                           title: str = Form("", max_length=60), db: DBSession = Depends(get_db),
                           user: User = Staff) -> dict:
    """Create a process narrative from a document (.docx, .xlsx or .txt): its text becomes the description."""
    from app.core.upload_policy import max_bytes, max_mb
    from app.domains.audit.rcm.models import Narrative
    from app.domains.audit.rcm.router import _audit_or_404, _ensure_open, _history, _objective

    obj = _objective(db, objective_id)
    audit = _audit_or_404(db, obj.audit_id)
    _ensure_open(audit)
    name = (file.filename or "narrative").strip()
    suffix = name[name.rfind("."):].lower() if "." in name else ""
    data = await file.read(max_bytes() + 1)
    if len(data) > max_bytes():
        raise field_error("file", f"the file is larger than {max_mb()} MB")
    try:
        text = _document_text(data, suffix).strip()
    except HTTPException:
        raise
    except Exception:  # noqa: BLE001 - a damaged document
        raise field_error("file", "the document could not be read") from None
    if not text:
        raise field_error("file", "the document has no text")
    count = db.execute(select(func.count()).select_from(Narrative).where(Narrative.objective_id == obj.id)
                       ).scalar_one()
    stem = name[: name.rfind(".")] if "." in name else name
    n = Narrative(objective_id=obj.id, ref=next_reference(db, "NAR"), title=(title.strip() or stem)[:60],
                  description=f"{text[:500_000]}\n\nImported from {name}.", position=count, custom_fields="{}",
                  created_by=user.id)
    db.add(n)
    db.flush()
    _history(db, audit.id, user, f"NARRATIVE {n.ref} IMPORTED from {name}", request)
    log_action(db, user_id=user.id, request=request, action=f"NARRATIVE_IMPORT id={n.id} audit={audit.id}")
    db.commit()
    return {"id": n.id, "ref": n.ref, "title": n.title, "audit_id": audit.id, "characters": len(text)}
