"""Statutory follow-up returns on audit recommendations.

* Quarterly follow-up report to the Audit Committee and the Internal Audit Agency (PFM Act 2016, Act 921
  s.83(5)(e); MoF Audit Recommendations Implementation and Follow-up Instructions s.6.1(3), s.7.5(3)): every
  recommendation with the finding, recommendation, management response, responsible officer, implementation
  status and the internal audit unit's follow-up comment.
* Annual statement on the implementation status of recommendations (Act 921 s.88(1)(b)), due by an
  admin-set date (30 June) of the following year.

Each submission is recorded in a register with a frozen copy of the rows, so what was sent can be reproduced.
"""

from __future__ import annotations

import json
from collections import Counter
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session as DBSession

from app.core.appconfig import cfg
from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import require_perm
from app.core.enums import CapaStatus, utcnow
from app.domains.issues.capa.models import Capa
from app.domains.issues.capa.status import IMPLEMENTATION_STATUSES
from app.domains.issues.followup.models import FollowupSubmission
from app.platform.users.models import User
from app.services.documents import add_table, docx_response, new_document, workbook, xlsx_response

router = APIRouter(prefix="/api/follow-up", tags=["follow-up reports"])
Reporter = Depends(require_perm("followup.report", "governance.view"))
Preparer = Depends(require_perm("followup.report"))

COLUMNS = [
    ("no", "No."), ("ref", "Ref"), ("source", "Source"), ("source_reference", "Source reference"),
    ("audit", "Audit / report"), ("finding", "Finding"), ("recommendation", "Recommendation"),
    ("management_response", "Management response / action plan"), ("responsible_officer", "Responsible officer"),
    ("timeframe", "Timeframe"), ("due_date", "Due date"), ("status", "Implementation status"),
    ("iau_comment", "Internal Audit follow-up comment"), ("days_overdue", "Days overdue"),
]


def quarter_bounds(year: int, quarter: int) -> tuple[datetime, datetime]:
    if quarter not in (1, 2, 3, 4):
        raise HTTPException(status_code=422, detail="Quarter must be 1 to 4")
    start = datetime(year, quarter * 3 - 2, 1)
    end = datetime(year + (quarter == 4), 1 if quarter == 4 else quarter * 3 + 1, 1)
    return start, end


def annual_due_date(year: int) -> datetime:
    text = str(cfg("annual_statement_due") or "06-30")
    try:
        month, day = (int(x) for x in text.split("-")[:2])
        return datetime(year + 1, month, day, 23, 59)
    except ValueError:
        return datetime(year + 1, 6, 30, 23, 59)


def _fmt(d: datetime | None) -> str:
    return f"{d:%d %b %Y}" if d else ""


def recommendation_rows(db: DBSession, start: datetime, end: datetime) -> list[dict]:
    """Recommendations open at any time in [start, end): raised before the end and not closed before the start."""
    from app.domains.audit.audits.models import Audit, Finding
    from app.domains.audit.reports.models import AuditReport

    capas = db.execute(select(Capa).where(Capa.created_at < end,
                                          or_(Capa.closed_at.is_(None), Capa.closed_at >= start))
                       .order_by(Capa.source_body, Capa.capa_ref)).scalars().all()
    now = utcnow()
    rows = []
    for i, c in enumerate(capas, 1):
        finding = db.get(Finding, c.source_id) if c.source_type == "finding" and c.source_id else None
        audit = db.get(Audit, finding.audit_id) if finding else None
        report = (db.execute(select(AuditReport).where(AuditReport.audit_id == audit.id)).scalars().first()
                  if audit else None)
        open_ = c.status != CapaStatus.Closed
        overdue = (now - c.due_date).days if open_ and c.due_date and c.due_date < now else 0
        rows.append({
            "no": i, "id": c.id, "ref": c.capa_ref, "source": c.source_body or "",
            "source_reference": c.source_reference or (report.report_ref if report else ""),
            "audit": (f"{report.report_ref} - " if report else "") + (audit.title if audit else ""),
            "finding": (finding.title + (f": {finding.condition}" if finding.condition else "")) if finding
            else (c.description or c.title),
            "recommendation": (finding.recommendation if finding and finding.recommendation else c.corrective_action)
            or c.title,
            "management_response": "\n".join(x for x in (
                finding.management_response if finding else "", c.corrective_action if finding else "",
                c.progress_note or "") if x),
            "responsible_officer": (finding.responsible_officer if finding and finding.responsible_officer
                                    else c.owner_email) or "",
            "timeframe": c.timeframe or "", "due_date": _fmt(c.due_date),
            "status": c.implementation_status or "Not Implemented",
            "iau_comment": "\n".join(x for x in (c.followup_comment or "", c.validation_comment or "",
                                                 f"Withdrawal proposed: {c.withdrawal_reason}"
                                                 if c.withdrawal_reason and open_ else "") if x),
            "days_overdue": overdue or "",
        })
    return rows


def summarise(rows: list[dict]) -> dict:
    by_status = Counter(r["status"] for r in rows)
    by_source: dict[str, Counter] = {}
    for r in rows:
        by_source.setdefault(r["source"] or "Unspecified", Counter())[r["status"]] += 1
    total = len(rows)
    done = by_status.get("Fully Implemented", 0)
    return {
        "total": total,
        "by_status": {s: by_status.get(s, 0) for s in IMPLEMENTATION_STATUSES},
        "by_source": {k: {s: v.get(s, 0) for s in IMPLEMENTATION_STATUSES} for k, v in by_source.items()},
        "overdue": sum(1 for r in rows if r["days_overdue"]),
        "implementation_rate": round(100 * done / (total - by_status.get("Withdrawn", 0)), 1)
        if total - by_status.get("Withdrawn", 0) else None,
    }


def _period(kind: str, year: int, quarter: int | None) -> tuple[str, datetime, datetime]:
    if kind == "quarterly":
        start, end = quarter_bounds(year, quarter or 0)
        return f"Q{quarter} {year}", start, end
    return f"Year {year}", datetime(year, 1, 1), datetime(year + 1, 1, 1)


def build(db: DBSession, kind: str, year: int, quarter: int | None) -> dict:
    if kind not in ("quarterly", "annual"):
        raise HTTPException(status_code=422, detail="Unknown report type")
    if not 2000 <= year <= utcnow().year + 1:
        raise HTTPException(status_code=422, detail="Choose a valid year")
    label, start, end = _period(kind, year, quarter)
    rows = recommendation_rows(db, start, end)
    return {
        "kind": kind, "year": year, "quarter": quarter, "period_label": label,
        "period_start": start, "period_end": end,
        "recipients": cfg("quarterly_report_recipients" if kind == "quarterly" else "annual_statement_recipients"),
        "due_date": annual_due_date(year) if kind == "annual" else None,
        "columns": [{"key": k, "label": v} for k, v in COLUMNS],
        "rows": rows, "summary": summarise(rows), "prepared_at": utcnow(),
    }


def _title(data: dict) -> str:
    if data["kind"] == "quarterly":
        return f"Quarterly report on the implementation of audit recommendations - {data['period_label']}"
    return (f"Statement on the implementation status of audit recommendations for the year {data['year']} "
            "(Public Financial Management Act, 2016 (Act 921), section 88(1)(b))")


def to_docx(data: dict) -> bytes:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    from app.services.documents import to_bytes

    doc = new_document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    doc.add_heading(_title(data), level=1)
    doc.add_paragraph(f"To: {', '.join(data['recipients'])}")
    if data.get("due_date"):
        doc.add_paragraph(f"Due: {_fmt(data['due_date'])}")
    s = data["summary"]
    doc.add_heading("1. Summary", level=2)
    add_table(doc, ["Status", "Number"], [[k, v] for k, v in s["by_status"].items()] + [["Total", s["total"]],
                                                                                      ["Overdue", s["overdue"]]])
    if s["implementation_rate"] is not None:
        doc.add_paragraph(f"Implementation rate (fully implemented, excluding withdrawn): {s['implementation_rate']}%")
    doc.add_heading("2. By source of recommendation", level=2)
    add_table(doc, ["Source", *IMPLEMENTATION_STATUSES],
              [[src, *counts.values()] for src, counts in s["by_source"].items()])
    doc.add_heading("3. Status of each recommendation", level=2)
    keys = [k for k, _ in COLUMNS]
    table = add_table(doc, [v for _, v in COLUMNS], [[r.get(k, "") for k in keys] for r in data["rows"]])
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7.5)
    doc.add_paragraph("Implementation status is shown as at the date of preparation.").runs[0].italic = True
    doc.add_heading("Sign-off", level=2)
    if data["kind"] == "annual":
        signers = ["Director-General (Principal Spending Officer)", "Chairperson, Audit Committee",
                   "Head of Internal Audit (GM Audit)"]
    else:
        signers = ["Head of Internal Audit (GM Audit)"]
    for who in signers:
        doc.add_paragraph(f"{who}\nName: ____________________   Signature: ____________   Date: __________")
    return to_bytes(doc)


def to_xlsx(data: dict) -> bytes:
    keys = [k for k, _ in COLUMNS]
    s = data["summary"]
    summary = [[k, v] for k, v in s["by_status"].items()] + [["Total", s["total"]], ["Overdue", s["overdue"]],
                                                              ["Implementation rate %", s["implementation_rate"]]]
    return workbook([
        ("Recommendations", [v for _, v in COLUMNS], [[r.get(k, "") for k in keys] for r in data["rows"]]),
        ("Summary", ["Item", "Value"], summary),
        ("By source", ["Source", *IMPLEMENTATION_STATUSES],
         [[src, *counts.values()] for src, counts in s["by_source"].items()]),
    ], title=_title(data))


def _name(data: dict, ext: str) -> str:
    stem = f"followup-Q{data['quarter']}-{data['year']}" if data["kind"] == "quarterly" else \
        f"annual-statement-{data['year']}"
    return f"{stem}.{ext}"


# ---------------------------------------------------------------- routes

@router.get("/quarterly")
def quarterly(year: int, quarter: int, db: DBSession = Depends(get_db), _: User = Reporter) -> dict:
    return build(db, "quarterly", year, quarter)


@router.get("/annual")
def annual(year: int, db: DBSession = Depends(get_db), _: User = Reporter) -> dict:
    return build(db, "annual", year, None)


@router.get("/{kind}.{fmt}")
def export(kind: str, fmt: str, year: int, quarter: int | None = None, db: DBSession = Depends(get_db),
           _: User = Reporter):  # noqa: ANN201
    data = build(db, kind, year, quarter)
    if fmt == "docx":
        return docx_response(to_docx(data), _name(data, "docx"))
    if fmt == "xlsx":
        return xlsx_response(to_xlsx(data), _name(data, "xlsx"))
    raise HTTPException(status_code=404, detail="Unknown format")


class SubmissionIn(BaseModel):
    kind: str
    year: int
    quarter: int | None = None
    recipients: list[str] = Field(min_length=1, max_length=20)
    submitted_at: datetime
    reference: str = Field(default="", max_length=120)


def _sub_out(s: FollowupSubmission) -> dict:
    return {"id": s.id, "kind": s.kind, "year": s.year, "quarter": s.quarter, "period_label": s.period_label,
            "recipients": json.loads(s.recipients or "[]"), "submitted_at": s.submitted_at, "reference": s.reference,
            "counts": json.loads(s.counts or "{}"), "recorded_by": s.recorded_by, "created_at": s.created_at}


@router.get("/submissions")
def submissions(db: DBSession = Depends(get_db), _: User = Reporter) -> list[dict]:
    rows = db.execute(select(FollowupSubmission).order_by(FollowupSubmission.submitted_at.desc())).scalars()
    return [_sub_out(s) for s in rows]


@router.post("/submissions", status_code=201)
def record_submission(payload: SubmissionIn, request: Request, db: DBSession = Depends(get_db),
                      user: User = Preparer) -> dict:
    data = build(db, payload.kind, payload.year, payload.quarter)
    submitted = payload.submitted_at.replace(tzinfo=None)
    if submitted > utcnow():
        raise HTTPException(status_code=422, detail="The submission date cannot be in the future")
    if submitted < data["period_start"]:
        raise HTTPException(status_code=422, detail="A report cannot be submitted before its period starts")
    sub = FollowupSubmission(kind=payload.kind, year=payload.year, quarter=payload.quarter,
                             period_label=data["period_label"], recipients=json.dumps(payload.recipients),
                             submitted_at=submitted, reference=payload.reference.strip() or None,
                             snapshot=json.dumps(data, default=str), counts=json.dumps(data["summary"]),
                             recorded_by=user.id)
    db.add(sub)
    db.flush()
    log_action(db, user_id=user.id, request=request,
               action=f"FOLLOWUP_SUBMITTED {payload.kind} {data['period_label']} to {', '.join(payload.recipients)}")
    db.commit()
    return _sub_out(sub)


@router.get("/submissions/{sub_id}/export.{fmt}")
def submission_export(sub_id: int, fmt: str, db: DBSession = Depends(get_db), _: User = Reporter):  # noqa: ANN201
    sub = db.get(FollowupSubmission, sub_id)
    if sub is None:
        raise HTTPException(status_code=404, detail="Submission not found")
    data = json.loads(sub.snapshot)
    data["recipients"] = json.loads(sub.recipients or "[]")
    if data.get("due_date"):
        data["due_date"] = datetime.fromisoformat(str(data["due_date"]))
    if fmt == "docx":
        return docx_response(to_docx(data), _name(data, "docx"))
    if fmt == "xlsx":
        return xlsx_response(to_xlsx(data), _name(data, "xlsx"))
    raise HTTPException(status_code=404, detail="Unknown format")
