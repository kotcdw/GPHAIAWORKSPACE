"""Excel import (org units, processes, findings) and Excel export of any list.

Imports are two-step: upload with ``commit=false`` to validate and preview, then
upload again with ``commit=true``. Nothing is written if any row has an error.
"""

from __future__ import annotations

import io
import re
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession

from app.core.access import ensure_can_view
from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import ensure_perm, require_active_user, require_module_tier, require_perm
from app.core.enums import AuditStatus, RiskRating, Role
from app.core.history import log_record_history
from app.domains.audit.audits.models import Audit, Finding
from app.platform.users.models import User

router = APIRouter(tags=["import / export"])
Staff = Depends(require_module_tier("import", Role.Admin, Role.Auditor))
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 20 * 1024 * 1024  # hard ceilings; the working limits are settings (Configuration > Data & attachments)
MAX_ROWS = 5000


def import_max_bytes() -> int:
    from app.core.appconfig import cfg

    return min(int(cfg("import_max_mb")) * 1024 * 1024, MAX_BYTES)


def import_max_rows() -> int:
    from app.core.appconfig import cfg

    return min(int(cfg("import_max_rows")), MAX_ROWS)
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

TEMPLATES: dict[str, list[tuple[str, bool, str]]] = {
    # column, required, help
    "orgunits": [
        ("Name", True, "Org unit name, e.g. Finance Department"),
        ("Type", True, "One of the Org Structure types (Configuration > Organizations), e.g. Department"),
        ("Parent", False, "Name or reference of the unit it sits under (an existing unit or one higher in the file)"),
        ("Address", False, ""),
    ],
    "processes": [
        ("Process", True, "Process name, e.g. Revenue collection"),
        ("Description", False, ""),
        ("Parent process", False, "Name of the process it belongs to (existing or higher in the file)"),
        ("Owner email", False, "Process owner"),
        ("Org units", False, "Org units that perform it, separated by ; (names or references)"),
    ],
    "findings": [
        ("Title", True, "Short statement of the issue"),
        ("Risk rating", True, "High, Medium or Low"),
        ("Description", False, ""),
        ("Criteria", False, "What should be (policy, law, standard)"),
        ("Condition", False, "What was found"),
        ("Cause", False, "Why it happened"),
        ("Effect", False, "Risk or impact"),
        ("Recommendation", False, ""),
        ("Owner email", False, "Management owner of the finding"),
    ],
    "users": [
        ("Full name", True, "e.g. Ama Mensah"),
        ("Username", True, "Letters, digits, '.', '_' or '-'; unique"),
        ("Email", True, "Work e-mail; unique"),
        ("Role", True, "Administrator, Manager, Auditor, Read-only User or Audit Client"),
        ("Approval level", False, "Managers only: Unit Head, Audit Manager, PIA or GM Audit (CAE)"),
        ("Workplace", False, "Name or reference of their unit in the Org Structure"),
    ],
    "remediation": [
        ("Title", True, "The recommendation or action, e.g. Reconcile vessel calls to invoices monthly"),
        ("Source", True, "Who made the recommendation: one of the configured sources, e.g. Internal Audit"),
        ("Source reference", False, "Report or letter reference, e.g. AG/2025/14"),
        ("Source year", False, "Year of the report, e.g. 2025"),
        ("Severity", True, "High, Medium or Low"),
        ("Owner email", False, "Responsible officer"),
        ("Due date", False, "Target date (a date cell or YYYY-MM-DD); past dates are kept as overdue"),
        ("Timeframe", False, "One of the configured timeframes"),
        ("Description", False, "Background or condition found"),
        ("Recommended action", False, "What management should do"),
    ],
    "risks": [
        ("Title", True, "Risk statement, e.g. Revenue leakage from unbilled vessel calls"),
        ("Category", False, "One of the configured risk categories"),
        ("Description", False, ""),
        ("Cause", False, ""),
        ("Consequence", False, ""),
        ("Owner email", False, "Risk owner (an existing account)"),
        ("Org unit", False, "Name or reference of the unit in the Org Structure"),
        ("Inherent likelihood", True, "1 (rare) to 5 (almost certain)"),
        ("Inherent impact", True, "1 (minor) to 5 (severe)"),
        ("Residual likelihood", False, "1 to 5; defaults to inherent"),
        ("Residual impact", False, "1 to 5; defaults to inherent"),
        ("Appetite", False, "One of the configured appetite levels"),
        ("Treatment", False, "One of the configured treatments"),
        ("Treatment plan", False, ""),
    ],
}

EXTRA_EXAMPLES = {
    "users": [["Kofi Mensah", "k.mensah", "k.mensah@gpha.gov.gh", "Auditor", "", "Finance Department"],
              ["Esi Boateng", "e.boateng", "e.boateng@gpha.gov.gh", "Manager", "Unit Head", "Tema Port"]],
    "remediation": [["Reconcile vessel calls to invoices monthly", "Internal Audit", "IA/2025/07", 2025, "High",
                     "finance.manager@gpha.gov.gh", "2026-12-31", "", "No reconciliation in 2025",
                     "Assign a reviewer and evidence the monthly reconciliation"]],
    "risks": [["Revenue leakage from unbilled vessel calls", "", "", "Manual billing", "Loss of revenue",
               "finance.manager@gpha.gov.gh", "Finance Department", 4, 4, 3, 3, "", "", ""]],
}


def record_run(db: DBSession, user: User, kind: str, filename: str | None, result: ImportResult,
               context: str | None = None) -> ImportResult:
    """Keep imports that were carried out and uploads that were rejected (with their first problems)."""
    import json

    from app.platform.dataio.models import ImportRun

    if not (result.errors or result.committed):
        return result  # a clean preview: nothing happened yet
    try:
        if result.errors:
            db.rollback()  # nothing from a rejected file is kept
        db.add(ImportRun(kind=kind, filename=(filename or "")[:255] or None,
                         outcome="imported" if result.committed else "rejected", total_rows=result.total_rows,
                         created=result.to_create if result.committed else 0,
                         updated=result.to_update if result.committed else 0, error_count=len(result.errors),
                         errors=json.dumps([e.model_dump() for e in result.errors[:200]]) if result.errors else None,
                         context=context, user_id=user.id, user_name=user.full_name or user.username))
        if result.errors:
            log_action(db, user_id=user.id, action=f"IMPORT_REJECTED kind={kind} problems={len(result.errors)}")
        db.commit()
    except Exception:  # noqa: BLE001 - the history must never break an import
        db.rollback()
    return result


def _workbook(title: str, headers: list[str], rows: list[list], notes: list[str] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.services.documents.excel import safe_cell, safe_row

    wb = Workbook()
    ws = wb.active
    ws.title = str(safe_cell(re.sub(r"[\[\]\*\?/\\:]", "", title))).lstrip("'")[:31] or "Sheet1"
    ws.append(safe_row(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A68")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(safe_row(row))
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
    ws.freeze_panes = "A2"
    if notes:
        help_ws = wb.create_sheet("Instructions")
        for line in notes:
            help_ws.append([safe_cell(line)])
        help_ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(content: bytes, filename: str) -> Response:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", filename).strip("-") or "export.xlsx"
    return Response(content=content, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{safe}"'})


# ------------------------------------------------------------------ export


class ExportIn(BaseModel):
    title: str = Field("Export", max_length=100)
    columns: list[str] = Field(..., min_length=1, max_length=60)
    rows: list[list[str | int | float | None]] = Field(default_factory=list, max_length=20000)


@router.post("/api/export.xlsx")
def export_xlsx(
    payload: ExportIn, request: Request, db: DBSession = Depends(get_db), user: User = Depends(require_active_user)
) -> Response:
    """Turn the rows a user is looking at into a formatted Excel workbook."""
    ensure_perm(user, "data.export")
    content = _workbook(payload.title, payload.columns, payload.rows)
    log_action(
        db, user_id=user.id, action=f"EXPORT_XLSX {payload.title[:60]} rows={len(payload.rows)}", request=request
    )
    db.commit()
    return _xlsx_response(content, f"{payload.title}-{date.today().isoformat()}.xlsx")


class ExportNoteIn(BaseModel):
    title: str = Field("Export", max_length=100)
    rows: int = Field(0, ge=0, le=10_000_000)
    format: str = Field("csv", pattern="^(csv|print)$")


@router.post("/api/export.log", status_code=204)
def export_logged(payload: ExportNoteIn, request: Request, db: DBSession = Depends(get_db),
                  user: User = Depends(require_active_user)) -> Response:
    """Record a CSV export made in the browser. The web app asks first and downloads only when this succeeds."""
    ensure_perm(user, "data.export")
    log_action(db, user_id=user.id, request=request, commit=True,
               action=f"EXPORT_{payload.format.upper()} {payload.title[:60]} rows={payload.rows}")
    return Response(status_code=204)


# ------------------------------------------------------------------ import


@router.get("/api/import/{kind}/template.xlsx")
def template(kind: str, user: User = Staff) -> Response:
    if kind not in TEMPLATES:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Unknown import type")
    cols = TEMPLATES[kind]
    example = {
        "orgunits": [["Finance Department", "Department", "Tema Port", ""],
                     ["Revenue Unit", "Unit", "Finance Department", ""]],
        "processes": [["Revenue collection", "Billing and collection of port dues", "", "revenue@gpha.gov.gh",
                       "Revenue Unit; Finance Department"]],
        "findings": [
            [
                "Vessel calls not reconciled to invoices",
                "High",
                "",
                "Billing policy s.4.2",
                "No monthly reconciliation in 2026",
                "No assigned reviewer",
                "Revenue leakage",
                "Assign and evidence a monthly reconciliation",
                "finance.manager@gpha.gov.gh",
            ]
        ],
    }.get(kind) or EXTRA_EXAMPLES[kind]
    notes = (
        [f"How to fill the {kind} import sheet", ""]
        + [f"{c} ({'required' if req else 'optional'}){': ' + h if h else ''}" for c, req, h in cols]
        + [
            "",
            "Keep the header row exactly as provided. One record per row. Delete the example row first.",
            f"Up to {import_max_rows()} rows per file.",
        ]
    )
    return _xlsx_response(
        _workbook(kind.title(), [c for c, _, _ in cols], example, notes), f"gpha-{kind}-import-template.xlsx"
    )


class RowIssue(BaseModel):
    row: int
    field: str
    message: str


class ImportResult(BaseModel):
    kind: str
    committed: bool
    total_rows: int
    to_create: int
    to_update: int
    errors: list[RowIssue]
    preview: list[dict]


def _read_rows(data: bytes, kind: str, *, db: DBSession | None = None, user: User | None = None,
               filename: str | None = None, context: str | None = None) -> tuple[list[dict], list[RowIssue]]:
    """Read the sheet. A file that cannot be used at all (too large, not a workbook, wrong columns) is refused, and
    the refusal is kept in the import history when ``db`` and ``user`` are given."""
    try:
        return _read_rows_unchecked(data, kind)
    except HTTPException as exc:
        if db is not None and user is not None:
            record_run(db, user, kind, filename, ImportResult(
                kind=kind, committed=False, total_rows=0, to_create=0, to_update=0, preview=[],
                errors=[RowIssue(row=0, field="file", message=str(exc.detail))]), context=context)
        raise


def _read_rows_unchecked(data: bytes, kind: str) -> tuple[list[dict], list[RowIssue]]:
    from openpyxl import load_workbook

    if len(data) > import_max_bytes():
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                            detail=f"The file is larger than {import_max_bytes() // (1024 * 1024)} MB")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="The file is not a valid Excel workbook (.xlsx)"
        ) from exc
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    expected = [c for c, _, _ in TEMPLATES[kind]]
    missing = [c for c, req, _ in TEMPLATES[kind] if req and c not in header]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Missing column(s): {', '.join(missing)}. Download the template to see the layout.",
        )
    idx = {c: header.index(c) for c in expected if c in header}
    rows, issues = [], []
    for n, raw in enumerate(it, start=2):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        rec = {c: (raw[i] if i < len(raw) else None) for c, i in idx.items()}
        rec = {c: (v.strip() if isinstance(v, str) else v) for c, v in rec.items()}
        rec["_row"] = n
        rows.append(rec)
        if len(rows) > import_max_rows():
            issues.append(RowIssue(row=n, field="file", message=f"More than {import_max_rows()} rows; split the file"))
            break
    if not rows:
        issues.append(RowIssue(row=1, field="file", message="The sheet has no data rows"))
    return rows, issues


def _required(rec: dict, cols: list[tuple[str, bool, str]], issues: list[RowIssue]) -> None:
    for c, req, _ in cols:
        if req and (rec.get(c) is None or str(rec.get(c)).strip() == ""):
            issues.append(RowIssue(row=rec["_row"], field=c, message="is required"))


def _rating(v) -> int | None:  # noqa: ANN001
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        return None
    return n if 1 <= n <= 5 and float(v) == n else None


@router.post("/api/import/orgunits", response_model=ImportResult)
async def import_org_units(
    request: Request,
    file: UploadFile = File(...),
    commit: bool = Query(False),
    db: DBSession = Depends(get_db),
    user: User = Depends(require_perm("system.reference", "method.library")),
) -> ImportResult:
    """Org Structure bulk add (create and update by name)."""
    from app.core.appconfig import cfg
    from app.core.refs import next_reference
    from app.platform.organizations.models import Organization

    rows, issues = _read_rows(await file.read(), "orgunits", db=db, user=user, filename=file.filename)
    cols = TEMPLATES["orgunits"]
    types = {t.lower(): t for t in (cfg("org_structure_categories") or [])}
    existing = list(db.execute(select(Organization)).scalars())
    by_name = {o.name.strip().lower(): o for o in existing}
    by_ref = {o.org_ref.upper(): o for o in existing}
    in_file: set[str] = set()
    plan: list[tuple[dict, object | None]] = []
    for rec in rows:
        _required(rec, cols, issues)
        r, name = rec["_row"], str(rec.get("Name") or "").strip()
        if len(name) > 255:
            issues.append(RowIssue(row=r, field="Name", message="must be 255 characters or fewer"))
        if rec.get("Type") and str(rec["Type"]).strip().lower() not in types:
            issues.append(RowIssue(row=r, field="Type", message=f"must be one of: {', '.join(types.values())}"))
        parent = str(rec.get("Parent") or "").strip()
        if parent and parent.lower() not in by_name and parent.upper() not in by_ref and parent.lower() not in in_file:
            issues.append(RowIssue(row=r, field="Parent", message="is not an existing unit or one higher in the file"))
        known = by_name.get(parent.lower()) or by_ref.get(parent.upper()) if parent else None
        if known is not None and not known.is_active and parent.lower() not in in_file:
            issues.append(RowIssue(row=r, field="Parent", message=f"{known.name} is deactivated: reactivate it first"))
        same = by_name.get(name.lower())
        if same is not None and not same.is_active:
            issues.append(RowIssue(row=r, field="Name", message="is a deactivated unit: reactivate it before updating"))
        if name.lower() in in_file:
            issues.append(RowIssue(row=r, field="Name", message="appears twice in the file"))
        in_file.add(name.lower())
        plan.append((rec, by_name.get(name.lower())))
    to_update = sum(1 for _, o in plan if o is not None)
    result = ImportResult(kind="orgunits", committed=False, total_rows=len(rows), to_create=len(plan) - to_update,
                          to_update=to_update, errors=issues,
                          preview=[{"row": rec["_row"], "name": rec.get("Name"), "type": rec.get("Type"),
                                    "parent": rec.get("Parent"), "action": "update" if o else "create"}
                                   for rec, o in plan[:50]])
    if not commit or issues:
        return record_run(db, user, "orgunits", file.filename, result)
    from app.core.enums import OrganizationStatus

    for rec, org in plan:
        name = str(rec["Name"]).strip()
        if org is None:
            org = Organization(org_ref=next_reference(db, "ORG"), name=name, status=OrganizationStatus.Approved,
                               notes="", created_by=user.id)
            db.add(org)
        org.category = types[str(rec["Type"]).strip().lower()]
        parent = str(rec.get("Parent") or "").strip()
        if parent:
            p = by_name.get(parent.lower()) or by_ref.get(parent.upper())
            db.flush()
            org.parent_id = p.id if p is not None and p is not org else org.parent_id
        if rec.get("Address"):
            org.address = str(rec["Address"])
        db.flush()
        by_name[name.lower()] = org
    log_action(db, user_id=user.id, action=f"IMPORT_ORG_UNITS created={result.to_create} updated={to_update}",
               request=request)
    db.commit()
    result.committed = True
    return record_run(db, user, "orgunits", file.filename, result)


@router.post("/api/import/processes", response_model=ImportResult)
async def import_processes(
    request: Request,
    file: UploadFile = File(...),
    commit: bool = Query(False),
    db: DBSession = Depends(get_db),
    user: User = Depends(require_perm("method.library", "system.reference")),
) -> ImportResult:
    """Process library bulk add (create and update by name), with the org units that perform each process."""
    from app.core.refs import next_reference
    from app.domains.assurance.universe.models import Process, ProcessOrgUnit
    from app.domains.assurance.universe.router import OrgTree

    rows, issues = _read_rows(await file.read(), "processes", db=db, user=user, filename=file.filename)
    cols = TEMPLATES["processes"]
    tree = OrgTree(db)
    org_by = {o.name.strip().lower(): o.id for o in tree.nodes.values()} | {o.org_ref.upper(): o.id
                                                                             for o in tree.nodes.values()}
    existing = {p.name.strip().lower(): p for p in db.execute(select(Process)).scalars()}
    in_file: set[str] = set()
    plan: list[tuple[dict, object | None, list[int]]] = []
    for rec in rows:
        _required(rec, cols, issues)
        r, name = rec["_row"], str(rec.get("Process") or "").strip()
        if name.lower() in in_file:
            issues.append(RowIssue(row=r, field="Process", message="appears twice in the file"))
        in_file.add(name.lower())
        parent = str(rec.get("Parent process") or "").strip().lower()
        if parent and parent not in existing and parent not in in_file:
            issues.append(RowIssue(row=r, field="Parent process", message="is not a process existing or higher up"))
        email = rec.get("Owner email")
        if email and not _EMAIL.match(str(email)):
            issues.append(RowIssue(row=r, field="Owner email", message="is not a valid e-mail address"))
        orgs = []
        for part in [x.strip() for x in str(rec.get("Org units") or "").split(";") if x.strip()]:
            oid = org_by.get(part.lower()) or org_by.get(part.upper())
            if oid is None:
                issues.append(RowIssue(row=r, field="Org units", message=f"'{part}' is not an org unit"))
            else:
                orgs.append(oid)
        plan.append((rec, existing.get(name.lower()), orgs))
    to_update = sum(1 for _, p, _ in plan if p is not None)
    result = ImportResult(kind="processes", committed=False, total_rows=len(rows), to_create=len(plan) - to_update,
                          to_update=to_update, errors=issues,
                          preview=[{"row": rec["_row"], "process": rec.get("Process"), "org_units": len(orgs),
                                    "action": "update" if p else "create"} for rec, p, orgs in plan[:50]])
    if not commit or issues:
        return record_run(db, user, "processes", file.filename, result)
    for rec, proc, orgs in plan:
        name = str(rec["Process"]).strip()
        if proc is None:
            proc = Process(ref=next_reference(db, "PROC"), name=name, status="Active", created_by=user.id)
            db.add(proc)
        proc.description = str(rec.get("Description") or "")
        proc.owner_email = str(rec["Owner email"]) if rec.get("Owner email") else None
        parent = str(rec.get("Parent process") or "").strip().lower()
        db.flush()
        if parent and parent in existing and existing[parent] is not proc:
            proc.parent_id = existing[parent].id
        if orgs:
            have = {o for (o,) in db.execute(select(ProcessOrgUnit.org_id).where(
                ProcessOrgUnit.process_id == proc.id)).all()}
            for oid in dict.fromkeys(orgs):
                if oid not in have:
                    db.add(ProcessOrgUnit(process_id=proc.id, org_id=oid))
        db.flush()
        existing[name.lower()] = proc
    log_action(db, user_id=user.id, action=f"IMPORT_PROCESSES created={result.to_create} updated={to_update}",
               request=request)
    db.commit()
    result.committed = True
    return record_run(db, user, "processes", file.filename, result)


@router.post("/api/audits/{audit_id}/import/findings", response_model=ImportResult)
async def import_findings(
    audit_id: int,
    request: Request,
    file: UploadFile = File(...),
    commit: bool = Query(False),
    db: DBSession = Depends(get_db),
    user: User = Staff,
) -> ImportResult:
    audit = db.get(Audit, audit_id)
    ensure_can_view(db, user, "audit", audit)
    if audit.status == AuditStatus.Closed:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="The audit is closed")
    rows, issues = _read_rows(await file.read(), "findings", db=db, user=user, filename=file.filename,
                              context=f"audit {audit_id}")
    cols = TEMPLATES["findings"]
    ratings = {r.value.lower(): r for r in RiskRating}
    for rec in rows:
        _required(rec, cols, issues)
        r = rec["_row"]
        if rec.get("Title") and len(str(rec["Title"])) > 255:
            issues.append(RowIssue(row=r, field="Title", message="must be 255 characters or fewer"))
        if rec.get("Risk rating") and str(rec["Risk rating"]).lower() not in ratings:
            issues.append(RowIssue(row=r, field="Risk rating", message="must be High, Medium or Low"))
        email = rec.get("Owner email")
        if email and not _EMAIL.match(str(email)):
            issues.append(RowIssue(row=r, field="Owner email", message="is not a valid e-mail address"))
    result = ImportResult(
        kind="findings",
        committed=False,
        total_rows=len(rows),
        to_create=len(rows),
        to_update=0,
        errors=issues,
        preview=[
            {"row": rec["_row"], "title": rec.get("Title"), "rating": rec.get("Risk rating"), "action": "create"}
            for rec in rows[:50]
        ],
    )
    if not commit or issues:
        return record_run(db, user, "findings", file.filename, result, context=f"audit {audit_id}")
    for rec in rows:
        f = Finding(
            audit_id=audit_id,
            title=str(rec["Title"]),
            description=str(rec.get("Description") or ""),
            risk_rating=ratings[str(rec["Risk rating"]).lower()],
            assigned_owner_email=str(rec["Owner email"]) if rec.get("Owner email") else None,
            criteria=str(rec.get("Criteria") or ""),
            condition=str(rec.get("Condition") or ""),
            cause=str(rec.get("Cause") or ""),
            effect=str(rec.get("Effect") or ""),
            recommendation=str(rec.get("Recommendation") or ""),
            created_by=user.id,
        )
        db.add(f)
    db.flush()
    log_action(db, user_id=user.id, action=f"IMPORT_FINDINGS audit={audit_id} rows={len(rows)}", request=request)
    log_record_history(
        db,
        record_type="audit",
        record_id=audit_id,
        user_id=user.id,
        action=f"IMPORTED {len(rows)} FINDING(S) FROM EXCEL",
        request=request,
    )
    db.commit()
    result.committed = True
    return record_run(db, user, "findings", file.filename, result, context=f"audit {audit_id}")
