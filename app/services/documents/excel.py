"""Excel workbooks for registers and statutory follow-up returns (same header styling everywhere)."""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from datetime import datetime
from urllib.parse import quote

from fastapi.responses import Response

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(sheets: Iterable[tuple[str, Sequence[str], Iterable[Sequence]]], *, title: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.core.appconfig import cfg

    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name[:31])
        start = 1
        if title:
            ws.cell(row=1, column=1, value=f"{cfg('org_name')} - {cfg('department_name')}").font = Font(bold=True)
            ws.cell(row=2, column=1, value=title).font = Font(bold=True, size=12)
            ws.cell(row=3, column=1, value=f"{cfg('document_classification')}  ·  generated "
                                            f"{datetime.now():%d %b %Y %H:%M}").font = Font(italic=True, size=9)
            start = 5
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=start, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths = [len(str(h)) for h in headers]
        r = start
        for r, row in enumerate(rows, start + 1):
            for col, value in enumerate(row, 1):
                if isinstance(value, datetime):
                    value = value.strftime("%d %b %Y")
                cell = ws.cell(row=r, column=col, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                widths[col - 1] = max(widths[col - 1], min(len(str(value or "")), 60))
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = max(10, min(w + 2, 60))
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        ws.auto_filter.ref = f"A{start}:{get_column_letter(max(len(headers), 1))}{max(r, start)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(data: bytes, filename: str) -> Response:
    return Response(content=data, media_type=XLSX_MEDIA_TYPE,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
