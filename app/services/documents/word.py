"""Word documents with the GPHA letterhead and classification footer."""

from __future__ import annotations

import io
from urllib.parse import quote

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi.responses import Response

from app.core.appconfig import cfg
from app.core.enums import utcnow

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def new_document(*, letterhead: bool = True) -> DocxDocument:
    """A blank GPHA document: Calibri 11, organisation / department letterhead, classification footer."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if letterhead:
        head = doc.add_paragraph()
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = head.add_run(f"{cfg('org_name').upper()}\n{cfg('department_name').upper()}")
        run.bold = True
        run.font.size = Pt(13)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = cfg("document_classification").strip()
    frun = footer.add_run(" · ".join(x for x in (label, f"Generated {utcnow():%d %B %Y} from the GPHA Internal "
                                                          "Audit Workspace") if x))
    frun.font.size = Pt(8)
    return doc


def add_table(doc: DocxDocument, headers: list[str], rows: list[list], style: str = "Light Grid Accent 1"):  # noqa: ANN201
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = "" if text is None else str(text)
    return table


def to_bytes(doc: DocxDocument) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_response(doc: DocxDocument | bytes, filename: str) -> Response:
    """Download response for a Word document (UTF-8 safe file name)."""
    content = doc if isinstance(doc, bytes) else to_bytes(doc)
    ascii_name = filename.encode("ascii", "replace").decode().replace('"', "'")
    return Response(content=content, media_type=DOCX_MEDIA_TYPE, headers={
        "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"})
