"""Export a presentation (each slide) or a dashboard (each tab) to PowerPoint or PDF.

Every report widget is run now with the viewer's own access, as when the item is opened, so an export never holds
data the viewer could not see. Charts become native PowerPoint charts; tables keep their first rows; text, images
and videos are placed where the author put them on the 12-column canvas.
"""

from __future__ import annotations

import base64
import io
import json
from typing import Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session as DBSession

from app.core.audit import log_action
from app.core.database import get_db
from app.core.deps import require_active_user
from app.core.permissions import has_perm
from app.platform.users.models import User
from app.services.reporting.router import _item, _run_widgets
from app.services.reporting.sharing import ensure_staff

router = APIRouter(prefix="/api/reporting", tags=["reporting"])

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_TABLE_ROWS = 15
MAX_TABLE_COLS = 8
# Theme: (background, text, accent) as RGB hex.
THEMES = {"Default": ("FFFFFF", "243247", "0B5D7B"), "Harbour": ("F3F8FB", "243247", "0B5D7B"),
          "Slate": ("20283D", "EEF2F7", "7FB8CE"), "High contrast": ("FFFFFF", "000000", "000000")}
SLIDE_W, SLIDE_H = 13.333, 7.5  # inches (16:9)
LEFT, TOP, GAP = 0.4, 1.15, 0.08


def _pages(db: DBSession, user: User, item_id: int) -> tuple[object, str, list[dict]]:
    it = _item(db, user, item_id)
    if it.kind == "report":
        raise HTTPException(status_code=422, detail="Export a report to Excel or CSV")
    d = json.loads(it.definition or "{}")
    pages = d.get("slides") if it.kind == "presentation" else d.get("tabs")
    if not pages:
        raise HTTPException(status_code=422, detail="There is nothing to export yet")
    ran = [{"name": p.get("name", ""), "widgets": _run_widgets(db, user, p.get("widgets", []), {})} for p in pages]
    return it, d.get("theme", "Default") if it.kind == "presentation" else "Default", ran


def _grid(widgets: list[dict]) -> tuple[float, float]:
    """Column width and row height (inches) so every widget fits on the page."""
    rows = max([w.get("y", 0) + w.get("h", 4) for w in widgets] or [1])
    width = SLIDE_W - 2 * LEFT
    height = SLIDE_H - TOP - 0.35
    return width / 12, min(height / max(rows, 1), 0.55)


def _box(w: dict, cw: float, rh: float) -> tuple[float, float, float, float]:
    x, y = w.get("x", 0), w.get("y", 0)
    span = min(w.get("w", 6), 12 - x)
    return LEFT + x * cw, TOP + y * rh, span * cw - GAP, w.get("h", 4) * rh - GAP


def _table_data(res: dict) -> tuple[list[str], list[list[str]]]:
    ct = res.get("crosstab")
    if ct:
        labels = {c["key"]: c["label"] for c in res["columns"]}
        head = [f"{labels.get(ct['row_key'], ct['row_key'])} / {labels.get(ct['column_key'], ct['column_key'])}",
                *ct["columns"], "Total"]
        body = [[r["value"], *map(str, r["cells"]), str(r["total"])] for r in ct["rows"]]
        return head[:MAX_TABLE_COLS], [r[:MAX_TABLE_COLS] for r in body][:MAX_TABLE_ROWS] + \
            [["Total", *map(str, ct["totals"]), str(ct["grand"])][:MAX_TABLE_COLS]]
    cols = res["columns"][:MAX_TABLE_COLS]
    body = [["" if r.get(c["key"]) is None else str(r.get(c["key"])) for c in cols]
            for r in res["rows"][:MAX_TABLE_ROWS]]
    return [c["label"] for c in cols], body


def _image_bytes(data_uri: str) -> bytes | None:
    try:
        return base64.b64decode(data_uri.split(",", 1)[1], validate=True)
    except (IndexError, ValueError):
        return None


def _video(db: DBSession, user: User, video_id: int | None):  # noqa: ANN202
    from app.services.reporting.media import can_see, media_path
    from app.services.reporting.models import ReportMedia

    m = db.get(ReportMedia, video_id) if video_id else None
    if m is None or not can_see(db, user, m) or not media_path(m).is_file():
        return None
    return m, media_path(m)


def _filename(name: str, ext: str) -> str:
    stem = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in name).strip() or "presentation"
    return f"{stem}.{ext}"


# ------------------------------------------------------------------ PowerPoint

def to_pptx(db: DBSession, user: User, name: str, theme: str, pages: list[dict]) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    bg, ink, accent = THEMES.get(theme, THEMES["Default"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    prs.core_properties.title = name

    def text_box(slide, x, y, w, h, text, size=12, bold=False, color=ink):  # noqa: ANN001, ANN202
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(str(text).split("\n") or [""]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = RGBColor.from_string(color)
        return tb

    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        if theme == "Harbour":
            bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor.from_string(accent)
            bar.line.fill.background()
        text_box(slide, LEFT, 0.3, SLIDE_W - 2 * LEFT, 0.7, page["name"], size=26, bold=True)
        cw, rh = _grid(page["widgets"])
        for w in sorted(page["widgets"], key=lambda w: (w.get("y", 0), w.get("x", 0))):
            x, y, bw, bh = _box(w, cw, rh)
            kind = w.get("type")
            if kind == "filter":
                continue
            title = w.get("title") or ""
            if title and kind != "text":
                text_box(slide, x, y, bw, 0.35, title, size=12, bold=True)
                y, bh = y + 0.38, max(bh - 0.38, 0.4)
            if kind == "text":
                text_box(slide, x, y, bw, bh, w.get("text", ""), size=14)
            elif kind == "image":
                data = _image_bytes(w.get("image", ""))
                if data:
                    from PIL import Image

                    with Image.open(io.BytesIO(data)) as im:
                        ratio = im.width / max(im.height, 1)
                    pw, ph = (bw, bw / ratio) if bw / ratio <= bh else (bh * ratio, bh)
                    slide.shapes.add_picture(io.BytesIO(data), Inches(x), Inches(y), Inches(pw), Inches(ph))
            elif kind == "video":
                found = _video(db, user, w.get("video_id"))
                if found:
                    m, path = found
                    slide.shapes.add_movie(str(path), Inches(x), Inches(y), Inches(bw), Inches(bh),
                                           mime_type=m.content_type)
            elif kind == "report":
                if w.get("error"):
                    text_box(slide, x, y, bw, bh, w["error"], size=11)
                    continue
                res = w["result"]
                if res.get("chart") and res.get("groups") and w.get("visualization", "both") != "table":
                    data = CategoryChartData()
                    groups = res["groups"][:20]
                    data.categories = [g["value"] for g in groups]
                    data.add_series("Records", [g["count"] for g in groups])
                    kind_ = {"pie": XL_CHART_TYPE.PIE, "line": XL_CHART_TYPE.LINE}.get(res["chart"],
                                                                                     XL_CHART_TYPE.BAR_CLUSTERED)
                    chart = slide.shapes.add_chart(kind_, Inches(x), Inches(y), Inches(bw), Inches(bh), data).chart
                    chart.has_legend = res["chart"] == "pie"
                    if chart.has_legend:
                        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.RIGHT, False
                    chart.font.size = Pt(10)
                    chart.font.color.rgb = RGBColor.from_string(ink)
                else:
                    head, body = _table_data(res)
                    rows = len(body) + 1
                    shape = slide.shapes.add_table(rows, len(head), Inches(x), Inches(y), Inches(bw),
                                                   Inches(min(bh, 0.3 * rows)))
                    table = shape.table
                    for c, label in enumerate(head):
                        table.cell(0, c).text = str(label)
                    for r, row in enumerate(body, start=1):
                        for c, val in enumerate(row):
                            table.cell(r, c).text = val
                    for r in range(rows):
                        for c in range(len(head)):
                            for para in table.cell(r, c).text_frame.paragraphs:
                                for run in para.runs:
                                    run.font.size = Pt(9)
                    if res["total"] > len(res["rows"][:MAX_TABLE_ROWS]) and not res.get("crosstab"):
                        text_box(slide, x, y + min(bh, 0.3 * rows) + 0.02, bw, 0.3,
                                 f"{res['total']} records; first {MAX_TABLE_ROWS} shown", size=9)
        text_box(slide, LEFT, SLIDE_H - 0.35, SLIDE_W - 2 * LEFT, 0.3,
                 f"{name} · data run with {user.full_name or user.username}'s access", size=8)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF

_LATIN = str.maketrans({"…": "...", "–": "-", "—": "-", "≥": ">=", "≤": "<=", "‘": "'", "’": "'", "“": '"', "”": '"',
                        "·": "-", "›": ">", "×": "x", "•": "-"})


def _t(value: object) -> str:
    """Text for the PDF core fonts (Latin-1); other characters are replaced."""
    return str(value).translate(_LATIN).encode("latin-1", "replace").decode("latin-1")


def to_pdf(db: DBSession, user: User, name: str, theme: str, pages: list[dict]) -> bytes:
    from fpdf import FPDF

    bg, ink, accent = THEMES.get(theme, THEMES["Default"])
    rgb = lambda h: tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))  # noqa: E731
    mm = 25.4
    pdf = FPDF(orientation="L", unit="mm", format=(SLIDE_H * mm, SLIDE_W * mm))
    pdf.set_auto_page_break(False)
    pdf.set_title(_t(name))
    for page in pages:
        pdf.add_page()
        pdf.set_fill_color(*rgb(bg))
        pdf.rect(0, 0, SLIDE_W * mm, SLIDE_H * mm, style="F")
        if theme == "Harbour":
            pdf.set_fill_color(*rgb(accent))
            pdf.rect(0, 0, SLIDE_W * mm, 3, style="F")
        pdf.set_text_color(*rgb(ink))
        pdf.set_font("Helvetica", "B", 22)
        pdf.set_xy(LEFT * mm, 0.35 * mm)
        pdf.cell(w=(SLIDE_W - 2 * LEFT) * mm, h=14, text=_t(page["name"]))
        cw, rh = _grid(page["widgets"])
        for w in sorted(page["widgets"], key=lambda w: (w.get("y", 0), w.get("x", 0))):
            x, y, bw, bh = (v * mm for v in _box(w, cw, rh))
            kind = w.get("type")
            if kind == "filter":
                continue
            title = w.get("title") or ""
            if title and kind != "text":
                pdf.set_font("Helvetica", "B", 11)
                pdf.set_xy(x, y)
                pdf.cell(w=bw, h=6, text=_t(title)[:120])
                y, bh = y + 7, max(bh - 7, 8)
            if kind == "text":
                pdf.set_font("Helvetica", "", 12)
                pdf.set_xy(x, y)
                pdf.multi_cell(w=bw, h=6, text=_t(w.get("text", ""))[:3000])
            elif kind == "image":
                data = _image_bytes(w.get("image", ""))
                if data:
                    pdf.image(io.BytesIO(data), x=x, y=y, w=bw, h=bh, keep_aspect_ratio=True)
            elif kind == "video":
                found = _video(db, user, w.get("video_id"))
                pdf.set_font("Helvetica", "I", 10)
                pdf.set_xy(x, y)
                pdf.multi_cell(w=bw, h=5, text=_t(f"Video: {found[0].filename if found else 'not available'} "
                                                  f"(plays in the presentation and the PowerPoint export)"))
            elif kind == "report":
                if w.get("error"):
                    pdf.set_font("Helvetica", "I", 10)
                    pdf.set_xy(x, y)
                    pdf.multi_cell(w=bw, h=5, text=_t(w["error"]))
                    continue
                res = w["result"]
                if res.get("chart") and res.get("groups") and w.get("visualization", "both") != "table":
                    groups = res["groups"][:12]
                    top = max(g["count"] for g in groups) or 1
                    line = min(7.0, bh / max(len(groups), 1))
                    label_w = bw * 0.32
                    pdf.set_font("Helvetica", "", 9)
                    for i, g in enumerate(groups):
                        gy = y + i * line
                        pdf.set_xy(x, gy)
                        pdf.cell(w=label_w, h=line, text=_t(g["value"])[:40])
                        pdf.set_fill_color(*rgb(accent))
                        pdf.rect(x + label_w, gy + line * 0.2, max((bw - label_w - 12) * g["count"] / top, 0.5),
                                 line * 0.6, style="F")
                        pdf.set_xy(x + bw - 11, gy)
                        pdf.cell(w=11, h=line, text=str(g["count"]), align="R")
                else:
                    head, body = _table_data(res)
                    col_w = bw / max(len(head), 1)
                    line = 5.0
                    pdf.set_font("Helvetica", "B", 8)
                    pdf.set_xy(x, y)
                    for label in head:
                        pdf.cell(w=col_w, h=line, text=_t(label)[:28], border="B")
                    pdf.set_font("Helvetica", "", 8)
                    for r, row in enumerate(body, start=1):
                        if r * line + line > bh:
                            break
                        pdf.set_xy(x, y + r * line)
                        for val in row:
                            pdf.cell(w=col_w, h=line, text=_t(val)[:34])
        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(*rgb(ink))
        pdf.set_xy(LEFT * mm, (SLIDE_H - 0.3) * mm)
        pdf.cell(w=(SLIDE_W - 2 * LEFT) * mm, h=4,
                 text=_t(f"{name} - data run with {user.full_name or user.username}'s access"))
    return bytes(pdf.output())


@router.get("/items/{item_id}/export/{fmt}")
def export_presentation(item_id: int, request: Request, fmt: Literal["pptx", "pdf"],
                        download: bool = Query(True), db: DBSession = Depends(get_db),
                        user: User = Depends(require_active_user)) -> Response:
    """PowerPoint (.pptx) or PDF of a presentation (one page per slide) or a dashboard (one page per tab)."""
    ensure_staff(user)
    if not has_perm(user, "data.export"):
        raise HTTPException(status_code=403, detail="Your role does not allow exporting data")
    it, theme, pages = _pages(db, user, item_id)
    data = (to_pptx if fmt == "pptx" else to_pdf)(db, user, it.name, theme, pages)
    log_action(db, user_id=user.id, request=request,
               action=f"REPORTING_EXPORT {it.kind} id={it.id} format={fmt} pages={len(pages)}")
    db.commit()
    name = _filename(it.name, fmt)
    media = PPTX_TYPE if fmt == "pptx" else "application/pdf"
    disposition = "attachment" if download else "inline"
    return Response(content=data, media_type=media,
                    headers={"Content-Disposition": f"{disposition}; filename*=UTF-8''{quote(name)}"})
